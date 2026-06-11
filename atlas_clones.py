"""Карточки-клоны конкурентов по Atlas: один сайт + один и тот же компрессор по спекам
(серия+номер, кВт, бар, произв., FF, VSD, ресивер), но РАЗНЫЕ URL/имена (дубли по
категориям). Именно они дают «неоднозначные» в спек-матче. Группы ≥2 карточек."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from collections import defaultdict
from atlas_spec_review import load_comp

OUT="/home/user/Statistic/Atlas_clones_review.xlsx"

def build():
    cands=load_comp()
    groups=defaultdict(list)
    for c in cands:
        groups[(c["site"], c["sn"], c["kw"], c["bar"], c["fl"],
                c["ff"] or 0, c["vsd"] or 0, c["rv"])].append(c)
    clone_groups=[g for g in groups.values() if len(g)>=2]
    # сортировка: сайт, серия, по убыванию размера группы
    clone_groups.sort(key=lambda g:(g[0]["site"], g[0]["sn"], -len(g)))
    print(f"групп-клонов {len(clone_groups)} | карточек {sum(len(g) for g in clone_groups)}")

    wb=openpyxl.Workbook(); ws=wb.active; ws.title="клоны Atlas"
    blue=Font(color="0563C1", underline="single")
    strike=Font(color="C00000", underline="single", strike=True)
    bold=Font(bold=True, color="FFFFFF"); hfill=PatternFill("solid", fgColor="305496")
    band=PatternFill("solid", fgColor="EDEDED")     # чередование групп
    center=Alignment(horizontal="center", vertical="center", wrap_text=True)
    HDR=["Группа","Сайт","Серия","кВт","бар","произв л/мин","FF","VSD","ресивер",
         "Карточек","Карточка (имя → ссылка)","Цена","Снято"]
    ws.append(HDR)
    for ci in range(1,len(HDR)+1):
        cell=ws.cell(1,ci); cell.font=bold; cell.fill=hfill; cell.alignment=center
    r=1
    for gi,g in enumerate(clone_groups,1):
        g=sorted(g, key=lambda c:(c["price"] or 1e18))
        rep=g[0]
        for j,c in enumerate(g):
            r+=1
            if gi%2==0:
                for cc in range(1,len(HDR)+1): ws.cell(r,cc).fill=band
            if j==0:                                   # шапка группы — общие спеки
                ws.cell(r,1,gi); ws.cell(r,2,c["site"])
                ws.cell(r,3, f"{c['sn'][0].upper()}{c['sn'][1]:g}")
                ws.cell(r,4,c["kw"]); ws.cell(r,5,c["bar"]); ws.cell(r,6,c["fl"])
                ws.cell(r,7,"FF" if c["ff"] else ""); ws.cell(r,8,"VSD" if c["vsd"] else "")
                ws.cell(r,9,c["rv"] if c["rv"] else ""); ws.cell(r,10,len(g))
            nm=ws.cell(r,11, c["name"][:70]); nm.hyperlink=c["url"]
            nm.font=strike if c["status"]=="снято" else blue
            if c["price"]:
                pc=ws.cell(r,12,c["price"]); pc.number_format="# ##0"
            else:
                ws.cell(r,12, "снято" if c["status"]=="снято" else "По запросу")
            if c["status"]=="снято": ws.cell(r,13,"снято").font=strike
    widths=[7,18,8,6,6,12,5,5,9,9,72,12,8]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="C2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(HDR))}{r}"
    wb.save(OUT); print(f"-> {OUT}")

if __name__=="__main__":
    build()
