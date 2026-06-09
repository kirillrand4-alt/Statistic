"""Отдельный отчёт «Снятое → Замена». Основные отчёты/matcher НЕ меняет — только читает.
Источник снятия серий — LLM-справочник актуальности (web-search). Берём только уверенно
снятые серии; «под вопросом» игнорируем. Замены сверяем с рынком (есть ли у конкурентов, цена)."""
import openpyxl, csv, re
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from matcher import find_brand, domain
from build_review import load_all, COMPETITORS, STATUS

U="/root/.claude/uploads/62a19005-a7bf-569b-926e-b59b4a62600d/"
SERIES_FILE = U+"f0a49055-______________________.xlsx"
PROKO_CSV   = U+"e7171060-products_export_20260608.csv"
OUT         = "/home/user/Statistic/Snyatye_zameny_review.xlsx"

BRMAP={"abac":"abac","ariacom":"ariacom","atlas copco":"atlas","atmos":"atmos","dalgakiran":"dalgakiran",
"fiac":"fiac","fubag":"fubag","ingersoll rand":"ir","kraftmachine":"kraftmachine","kraftmann":"kraftmann",
"pneumatech":"pneumatech"}
# официальная замена -> (бренд замены, токен для поиска на рынке)
REPL={"FORMULA.I":("abac","genesis"),"APD-B":("ariacom","apdp"),"APD-V":("ariacom","apdp"),
"APD-S":("ariacom","apdp"),"XRYS":("atlas","xrvs"),"AHD":("ariacom","arhp"),"UP":("ir","up6")}
DOMFIX={"rostov.pnevmo-sklad.ru":"pnevmo-sklad.ru","novosibirsk.pnevmo-sklad.ru":"pnevmo-sklad.ru"}
def dm(u): d=domain(u).replace("www.",""); return DOMFIX.get(d,d)
def norm(s): return re.sub(r'[^a-zа-я0-9]','',str(s).lower())

def discontinued():
    ws=openpyxl.load_workbook(SERIES_FILE, read_only=True, data_only=True)["Лист1"]
    out=[]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[8] and ("снят" in str(r[8]).lower() or "архив" in str(r[8]).lower()):
            out.append((str(r[0]).strip(), str(r[1]).strip(),
                        str(r[14]).strip() if r[14] else "", str(r[13]).strip() if r[13] else "—"))
    return out

def market(info, rbrand, rtok):
    """по токену замены: (сколько у нас, сайтов-конкур., диапазон цен, пример-URL конкур.)"""
    prices=[]; sites=set(); ours=0; ex=None
    for u,(name,pr) in info.items():
        if find_brand(u)!=rbrand: continue
        if rtok not in norm(name) and rtok not in norm(u): continue
        d=dm(u)
        if d=="prokompressor.ru": ours+=1
        elif d in COMPETITORS and STATUS.get(u)!="снято":
            sites.add(d)
            p=min([x for x in pr if x is not None], default=None)
            if p: prices.append(p); ex=ex or u
    rng=f"{int(min(prices)):,}–{int(max(prices)):,}".replace(","," ") if prices else "—"
    return ours, len(sites), rng, ex

def our_products():
    prod=[]
    with open(PROKO_CSV, encoding="utf-8-sig", errors="replace") as fh:
        rd=csv.reader(fh, delimiter=";"); next(rd, None)
        for r in rd:
            if len(r)<3 or "prokompressor" not in r[1]: continue
            try: p=float(str(r[2]).replace(",",".").replace(" ","")) or None
            except: p=None
            prod.append((r[0].strip().replace("&quot;",'"'), r[1].strip(), find_brand(r[1]), p))
    return prod

def build():
    info=load_all(); disc=discontinued(); prod=our_products()
    mcache={k:market(info,*v) for k,v in REPL.items()}
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Снятое-Замена"
    HDR=["№","Бренд","Серия (снята)","Наш товар","Наша цена","Наша ссылка","Замена (офиц.)",
         "Замена у нас","Замена у конкур. (сайтов)","Цена замены","Ссылка на замену","Признаки снятия"]
    ws.append(HDR)
    blue=Font(color="0563C1",underline="single"); bold=Font(bold=True,color="FFFFFF")
    hf=PatternFill("solid",fgColor="305496"); warn=PatternFill("solid",fgColor="FFE699")
    for c in range(1,len(HDR)+1):
        cc=ws.cell(1,c); cc.font=bold; cc.fill=hf; cc.alignment=Alignment(horizontal="center",wrap_text=True)
    n=0
    for brand,series,priznaki,rep in disc:
        bc=BRMAP.get(brand.lower()); sn=norm(series)
        if not bc or len(sn)<3: continue
        rb=REPL.get(series); ours=nsites=0; rng="—"; ex=None
        if rb: ours,nsites,rng,ex=mcache[series]
        for nm,u,b,p in prod:
            if b!=bc or sn not in norm(nm): continue
            n+=1; r=n+1
            ws.cell(r,1,n); ws.cell(r,2,brand); ws.cell(r,3,series); ws.cell(r,4,nm)
            c5=ws.cell(r,5,p if p else "—")
            if p: c5.number_format="# ##0"
            l=ws.cell(r,6,"открыть"); l.hyperlink=u; l.font=blue
            ws.cell(r,7, rep if rep and rep!="None" else "—")
            ws.cell(r,8, ours if rb else "—"); ws.cell(r,9, nsites if rb else "—"); ws.cell(r,10, rng)
            if ex: lk=ws.cell(r,11,"замена"); lk.hyperlink=ex; lk.font=blue
            ws.cell(r,12, priznaki[:80])
            if rb and ours==0 and nsites>0: ws.cell(r,8).fill=warn   # у нас нет, у конкурентов есть
    widths=[5,12,16,44,11,8,30,10,14,16,9,46]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(HDR))}{n+1}"
    wb.save(OUT); print(f"-> {OUT}  ({n} строк)")

if __name__=="__main__":
    build()
