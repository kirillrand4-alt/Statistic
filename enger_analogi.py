"""Аналоги бренда Enger у конкурентов (6 сайтов) — СПЕК-матчинг (не по названию: у Enger
свои коды серий, на конкурентах их нет). Сцепка строгая: тип (форма винт/поршень + масло)
обязателен, мощность по номиналу, давление ±3%, производительность ±5%, частотник без
явного конфликта (молчание=совместимо). Тип подтверждаем явной формой ИЛИ маслом+произв.
Выход: xlsx, лист «Все аналоги» (строка=пара Энгер↔аналог, со столбцом «почему») +
лист «Сводка по моделям» (число аналогов, мин.цена конкурента, кто дешевле Энгера)."""
import re, sys
sys.path.insert(0,".")
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from brand_spec_review import load_ours_all, load_comp_all

OUT="/home/user/Statistic/Enger_analogi.xlsx"
NOMINAL=[0.55,0.75,1.1,1.5,2.2,3,4,5.5,7.5,11,15,18.5,22,30,37,45,55,75,90,110,132,160,
         200,250,315,355,400,450,500,560,630,710,800,900,1000,1250,1600]
def nominal(v): return min(NOMINAL, key=lambda x: abs(x-v)) if v else None
def g(v): return "" if v is None else (f"{v:g}" if isinstance(v,float) else str(v))
def yn(v): return "—" if v is None else ("да" if v else "нет")
def disp(b): return "Ingersoll Rand" if b=="ir" else ("Atlas Copco" if b=="atlas" else b.capitalize())

# форма (винт/поршень/спираль/центроб) из названия; у Enger ВД её часто нет -> None,
# тогда тип подтверждаем маслом+производительностью (см. analog_ok).
FORMS=[("винт","винтов"),("поршень","поршнев"),("спираль","спиральн"),("спираль","scroll"),
       ("центроб","центробежн"),("ротор","роторн")]
FORM_DISP={"винт":"винтовой","поршень":"поршневой","спираль":"спиральный",
           "центроб":"центробежный","ротор":"роторный"}
def form_of(name):
    s=(name or "").lower()
    for key,pat in FORMS:
        if pat in s: return key
    return None
def near(a,b,tol): return a is not None and b is not None and abs(a-b)<=tol*max(a,b)

def analog_ok(o,c):
    """Строгий аналог. Возврат True/False. Тип ОБЯЗАН быть подтверждён (формой или масло+произв)."""
    if not o.get("kw") or not c.get("kw"): return False
    if nominal(o["kw"])!=nominal(c["kw"]): return False           # мощность по номиналу
    if not o.get("bar") or not c.get("bar"): return False         # давление знаем с обеих
    if not near(o["bar"],c["bar"],0.03): return False             # ±3%
    if o.get("oil") and c.get("oil") and o["oil"]!=c["oil"]: return False   # масло≠безмасло — режем
    if o.get("fl") and c.get("fl") and not near(o["fl"],c["fl"],0.05): return False  # ±5% если оба
    ov,cv=o.get("vsd"),c.get("vsd")
    if ov is not None and cv is not None and ov!=cv: return False  # частотник: явный конфликт
    of,cf=form_of(o["name"]),form_of(c["name"])
    if of and cf:                                                 # форма явна с обеих — должна совпасть
        return of==cf
    # формы нет с одной/двух сторон -> тип подтверждаем маслом + производительностью
    return bool(o.get("oil") and c.get("oil") and o["oil"]==c["oil"]
                and o.get("fl") and c.get("fl") and near(o["fl"],c["fl"],0.05))

def why(o,c):                                                     # «почему аналог» — совпавшее с числами
    p=[f"{g(nominal(o['kw']))} кВт"]
    p.append(f"{g(o['bar'])}≈{g(c['bar'])} бар")
    if o.get("fl") and c.get("fl"): p.append(f"{g(o['fl'])}≈{g(c['fl'])} л/мин")
    of,cf=form_of(o["name"]),form_of(c["name"])
    if of and cf: p.append(FORM_DISP.get(of,of))
    if o.get("oil") and c.get("oil"): p.append(f"{o['oil']}={c['oil']}")
    if o.get("vsd") is not None and c.get("vsd") is not None: p.append(f"частотник {yn(o['vsd'])}={yn(c['vsd'])}")
    return " · ".join(p)

def build():
    ours=load_ours_all(); cands=load_comp_all()
    eng=ours.get("enger",[])
    # дедуп моделей Энгера: одна строка на (серия,кВт,бар,произв,частотник,ресивер); цена — мин.
    emap={}
    for o in eng:
        if not o.get("kw") or not o.get("bar"): continue
        k=(o["sn"], round(o["kw"],1), round(o["bar"],1),
           round(o["fl"]) if o.get("fl") else None, o.get("vsd"), o.get("rv"))
        cur=emap.get(k)
        if cur is None or (o.get("price") and (not cur.get("price") or o["price"]<cur["price"])):
            emap[k]=o
    engmodels=list(emap.values())

    # индекс конкурентов по номиналу мощности (быстрый прунинг), бренд кладём в карточку
    comp_by_kw=defaultdict(list)
    for b,lst in cands.items():
        if b=="enger": continue                                  # сам бренд — не аналог
        for c in lst:
            if c.get("kw"): comp_by_kw[nominal(c["kw"])].append((b,c))

    sheet1=[]; summary=[]
    for o in engmodels:
        hits=[(b,c) for (b,c) in comp_by_kw.get(nominal(o["kw"]),[]) if analog_ok(o,c)]
        # группируем предложения в МОДЕЛИ-аналоги: (бренд,серия,кВт,бар,частотник,ресивер)
        groups=defaultdict(list)
        for b,c in hits:
            gk=(b, c["sn"], round(c["kw"]) if c.get("kw") else None,
                round(c["bar"]) if c.get("bar") else None, c.get("vsd"), c.get("rv"))
            groups[gk].append(c)
        analogs=[]
        for gk,offers in groups.items():
            b=gk[0]
            priced=[x for x in offers if x.get("price") and x.get("status")!="снято"]
            best=min(priced,key=lambda x:x["price"]) if priced else offers[0]
            name=max(offers,key=lambda x:len(x.get("name") or "")).get("name") or ""
            allp=sorted({(x["site"],int(x["price"])) for x in offers if x.get("price")},
                        key=lambda t:t[1])
            analogs.append(dict(brand=b, name=name, why=why(o,best),
                                price=best.get("price"), site=best["site"], url=best["url"],
                                offers="; ".join(f"{s}: {p:,}".replace(","," ") for s,p in allp)))
        if not analogs: continue
        analogs.sort(key=lambda a:(a["price"] is None, a["price"] or 0))
        ofm=form_of(o["name"]); otype=FORM_DISP.get(ofm,"") or (("безмасляный " if o.get("oil")=="безмасло" else "")+ "ВД" if o["bar"]>=20 else (o.get("oil") or ""))
        eprice=o.get("price")
        for a in analogs:                                        # лист 1: пара на строку
            d=("" if (eprice is None or a["price"] is None) else round((a["price"]-eprice)/eprice*100))
            sheet1.append([o["name"], otype, o.get("kw"), o.get("bar"), o.get("fl"), yn(o.get("vsd")),
                           eprice, disp(a["brand"]), a["name"], a["why"], a["price"], a["site"],
                           d, a["offers"], o.get("url"), a["url"]])
        # лист 2: сводка по модели
        pr=[a["price"] for a in analogs if a["price"]]
        mn=min(pr) if pr else None
        mnb=disp(next(a["brand"] for a in analogs if a["price"]==mn)) if mn else ""
        brands=sorted({disp(a["brand"]) for a in analogs})
        cheaper=("" if (mn is None or eprice is None) else
                 (f"да (-{round((eprice-mn)/eprice*100)}%)" if mn<eprice else "нет"))
        summary.append([o["name"], otype, o.get("kw"), o.get("bar"), o.get("fl"), yn(o.get("vsd")),
                        eprice, len(analogs), len(brands), mn, mnb, cheaper, ", ".join(brands)])

    _save(sheet1, summary)
    nmod=len({r[0] for r in sheet1})
    print(f"моделей Энгера обработано: {len(engmodels)} | с аналогами: {nmod} | строк-пар: {len(sheet1)}")
    print(f"-> {OUT}")

H1=["Модель Энгер","Тип","кВт","бар","л/мин","Частотник","Цена Энгер ₽","Бренд аналога",
    "Модель аналога","Почему аналог (совпавшие спеки)","Мин. цена аналога ₽","Сайт (мин.)",
    "Δ к Энгеру, %","Все предложения (сайт: цена)","Ссылка Энгер","Ссылка аналога"]
H2=["Модель Энгер","Тип","кВт","бар","л/мин","Частотник","Цена Энгер ₽","Аналогов",
    "Брендов","Мин. цена конкур. ₽","Бренд (мин.)","Дешевле Энгера?","Бренды-аналоги"]
def _save(sheet1, summary):
    wb=openpyxl.Workbook()
    hfill=PatternFill("solid",fgColor="305496"); bold=Font(bold=True,color="FFFFFF")
    ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
    def put(ws,HDR,data,pricecols):
        ws.append(HDR)
        for ci in range(1,len(HDR)+1):
            c=ws.cell(1,ci); c.font=bold; c.fill=hfill; c.alignment=ctr
        for r in data: ws.append(r)
        for col in pricecols:
            for rr in range(2,ws.max_row+1):
                cell=ws.cell(rr,col)
                if isinstance(cell.value,(int,float)): cell.number_format="# ##0"
        ws.freeze_panes="A2"
    ws1=wb.active; ws1.title="Все аналоги"; put(ws1,H1,sheet1,[7,11])
    ws2=wb.create_sheet("Сводка по моделям"); put(ws2,H2,summary,[7,10])
    wb.save(OUT)

if __name__=="__main__":
    build()
