"""Матчинг ВСЕХ брендов ТОЛЬКО с compressortyt.ru — один xlsx. Единый НОВЫЙ формат
парсера (prokompressor.ru scrape + compressortyt.ru scrape, одинаковые колонки:
brand,name,model,specs(JSON),price,product_url). Спеки: Мощность/Рабочее давление/
Производительность/Осушитель/IP; частотник/ресивер — из имени (text_flags). Матч —
боевой spec_match.match (серия+номер обязательны, спеки «не противоречат») + фильтры
ресивер/FF/охлажд/IP. Статусы: есть у обоих / только у нас / нет у нас (GAP по compressortyt)."""
import csv, sys, json, re, os
sys.path.insert(0,".")
csv.field_size_limit(10**8)
from collections import defaultdict, Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from matcher import brand_from_text, brand_of
from spec_match import (num, sane_kw, bar_value, bar_from_text, flow_value, text_flags,
                        is_compressor, match, receiver_filter, ff_filter, ip_filter,
                        cool_filter, ip_class)
from brand_spec_review import ser_of, oil_of, suffix_flags

U="/root/.claude/uploads/62a19005-a7bf-569b-926e-b59b4a62600d/"
OURS=U+"2c8f1bfb-prices_20260710_110047.csv"          # наш prokompressor (со специями)
COMP=U+"803cc5a8-prices_20260710_110407.csv"          # compressortyt (специи заполнит пользователь)
OUT="/home/user/Statistic/Matching_compressortyt.xlsx"

def dom(u): return re.sub(r"^https?://","",u or "").split("/")[0].replace("www.","")
def spec_get(sp,*subs):
    for k,v in sp.items():
        kl=k.lower()
        if any(su in kl for su in subs): return v
    return None
def parse_row(r):
    name=(r.get("name") or "").strip(); model=(r.get("model") or "").strip()
    url=(r.get("product_url") or "").strip()
    if not is_compressor(name+" "+model): return None
    bstr=(r.get("brand") or "").strip()
    b=brand_from_text(bstr) or brand_from_text(name) or brand_of(url,name)
    if not b: return None
    text=name+" "+model
    sn=ser_of(text, b)
    if not sn: return None
    sp={}
    s=(r.get("specs") or "").strip()
    if s and s!="{}":
        try: sp=json.loads(s)
        except: sp={}
    kw=sane_kw(num(spec_get(sp,"мощн")))
    bar=bar_value(spec_get(sp,"давлен")) or bar_from_text(text)
    fl=flow_value(spec_get(sp,"производ","л/мин"),"л/мин")
    ff,vsd,rv=text_flags(text)
    ff,rv=suffix_flags(name,b,ff,rv)
    dry=str(spec_get(sp,"осушит") or "").strip().lower()
    if ff is None and dry in ("да","yes","есть"): ff=1
    elif ff is None and dry=="нет": pass
    ipv=spec_get(sp,"ip ","защит")
    ip=ip_class(text) or (ip_class(str(ipv)) if ipv else None)
    price=None
    try:
        pv=float(str(r.get("price","")).replace(",",".").replace(" ",""))
        if 100<=pv<=50_000_000: price=pv
    except: pass
    return dict(sn=sn,kw=kw,bar=bar,fl=fl,oil=oil_of(name),ff=ff,vsd=vsd,rv=rv,ip=ip,
                cool=None,dr=None,name=name,url=url,price=price,site=dom(url),brand=b,bdisp=bstr or b.capitalize())

def load(path):
    by=defaultdict(list)
    for r in csv.DictReader(open(path,encoding="utf-8-sig",errors="replace")):
        c=parse_row(r)
        if c: by[c["brand"]].append(c)
    return by

HEAD=["бренд","статус","наш товар","наши спеки","цена наша","карточка compressortyt",
      "спеки конкурента","цена конкур","наша ссылка","ссылка конкурента"]
ST={"есть у обоих":0,"только у нас":1,"нет у нас (GAP)":2}
def spc(o):
    p=[]
    if o.get("kw") is not None: p.append(f"{o['kw']:g}кВт")
    if o.get("bar") is not None: p.append(f"{o['bar']:g}бар")
    if o.get("fl") is not None: p.append(f"{o['fl']:g}л/мин")
    if o.get("ff"): p.append("FF");
    if o.get("vsd"): p.append("VSD")
    if o.get("rv"): p.append("ресив")
    return " ".join(p)
def rowdata(b,st,o,c):
    return [o.get("bdisp") if o else (c.get("bdisp") if c else b), st,
            o.get("name","") if o else "", spc(o) if o else "", o.get("price") if o else None,
            c.get("name","") if c else "", spc(c) if c else "", c.get("price") if c else None,
            o.get("url","") if o else "", c.get("url","") if c else ""]

def build():
    ours=load(OURS); comp=load(COMP)
    ncomp_spec=sum(1 for lst in comp.values() for c in lst if c["kw"] or c["bar"] or c["fl"])
    ncomp=sum(len(v) for v in comp.values())
    rows=[]; matched=set()
    for b in set(ours)|set(comp):
        by=defaultdict(list)
        for c in comp.get(b,[]): by[c["sn"]].append(c)
        for o in ours.get(b,[]):
            m=receiver_filter(o.get("rv"), ff_filter(o.get("ff"),
                cool_filter(o.get("cool"), ip_filter(o.get("ip"), match(o, by.get(o["sn"], []))))))
            if m:
                for c in m:
                    matched.add(c["url"]); rows.append(rowdata(b,"есть у обоих",o,c))
            else:
                rows.append(rowdata(b,"только у нас",o,None))
    for b,lst in comp.items():
        seen=set()
        for c in lst:
            if c["url"] in matched or c["url"] in seen: continue
            seen.add(c["url"]); rows.append(rowdata(b,"нет у нас (GAP)",None,c))
    rows.sort(key=lambda r:(r[0] or "", ST.get(r[1],9), r[2] or r[5]))

    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Матчинг compressortyt"
    hf=PatternFill("solid",fgColor="305496"); bd=Font(bold=True,color="FFFFFF")
    ctr=Alignment(horizontal="center",vertical="center",wrap_text=True); blue=Font(color="0563C1",underline="single")
    ws.append(HEAD)
    for ci in range(1,len(HEAD)+1):
        c=ws.cell(1,ci); c.font=bd; c.fill=hf; c.alignment=ctr
    for r in rows:
        ws.append(r); rr=ws.max_row
        for col in (5,8):
            if ws.cell(rr,col).value is not None: ws.cell(rr,col).number_format="# ##0"
        for col in (9,10):
            u=ws.cell(rr,col).value
            if u: ws.cell(rr,col).hyperlink=u; ws.cell(rr,col).font=blue
    ws.freeze_panes="A2"; ws.column_dimensions["C"].width=46; ws.column_dimensions["F"].width=46
    wb.save(OUT)
    cnt=Counter(r[1] for r in rows)
    print(f"наших {sum(len(v) for v in ours.values())} | compressortyt {ncomp} (со спеками {ncomp_spec})")
    print(f"есть у обоих: {cnt['есть у обоих']} | только у нас: {cnt['только у нас']} | GAP: {cnt['нет у нас (GAP)']}")
    print(f"-> {OUT}")

if __name__=="__main__":
    build()
