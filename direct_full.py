"""Полная раскатка кампаний Я.Директ из prokompressor.ru.csv: компрессоры (без Enger),
кампания = НОМИНАЛьная мощность (15.1->15), файл на кампанию (xlsx), всё в zip.
Фразы (6 шаблонов + autotargeting, кросс-минусовка, лимит 7 слов):
  название | [название] | компрессор название | <тип> компрессор название | название купить | название цена
Минус на группу[60]: др.давления + др.IP + частотник(у группы без частотника). Ставка
autotargeting=50, остальные пусто. Бренд atlas->Atlas Copco. Цены из прайса 26.06."""
import csv, sys, re, os, zipfile
sys.path.insert(0,".")
from collections import defaultdict
import openpyxl
from matcher import brand_from_text, find_brand
from spec_match import is_compressor, num, bar_value

SRC="/root/.claude/uploads/62a19005-a7bf-569b-926e-b59b4a62600d/426094e6-prokompressor.ru.csv"
TPL="/root/.claude/uploads/62a19005-a7bf-569b-926e-b59b4a62600d/6dbcf451-prices_20260626_041449.csv"
NP ="/root/.claude/uploads/62a19005-a7bf-569b-926e-b59b4a62600d/8871701e-prices_20260626_041449.csv"
OUTDIR="/home/user/Statistic/direct_campaigns"; ZIP="/home/user/Statistic/Direct_campaigns_FULL.zip"
NOMINAL=[0.55,0.75,1.1,1.5,2.2,3,4,5.5,7.5,11,15,18.5,22,30,37,45,55,75,90,110,132,160,
         200,250,315,355,400,450,500,560,630,710,800,900,1000,1250,1600]
_BD={"atlas":"Atlas Copco","ir":"Ingersoll Rand"}
def disp(b): return "" if not b else _BD.get(b, b.capitalize())
def nrm(u): u=(u or "").strip().lower(); return re.sub(r"^https?://","",u).replace("www.","").rstrip("/")
def trimw(s,n):
    if len(s)<=n: return s
    cut=s[:n]; m=max(cut.rfind(" "),cut.rfind("-"))
    return (cut[:m] if m>n*0.6 else cut).rstrip(" -")
def nominal(v): return min(NOMINAL, key=lambda x: abs(x-v))
TYPE_RE=re.compile(r"\(\s*бустер\s*\)|\b(винтов\w*|поршнев\w*|спиральн\w*|двухступенчат\w*|"
                   r"дизельн\w*|передвижн\w*|роторн\w*|центробежн\w*|безмасл\w*|масл\w*|"
                   r"компрессор\w*|дожимн\w*|бустер|высокого|низкого|давлени\w*)\b", re.I)
def product_type(nm):
    s=nm.lower(); return "бустер" if ("бустер" in s or "дожимн" in s) else "компрессор"
def type_adj(nm):
    s=nm.lower()
    for k,w in (("винтов","винтовой"),("поршнев","поршневой"),("спиральн","спиральный"),
                ("роторн","роторный"),("центробежн","центробежный")):
        if k in s: return w
    return "винтовой"
def headline(typ, model):
    for h in (f"Купить {typ} {model} по спец цене", f"Купить {typ} {model}",
              f"Купить {model} по спец цене", f"Купить {model}"):
        if len(h)<=56: return h
    return trimw(f"Купить {model}",56)
def clean_model(s):
    s=re.sub(r"ATLAS\s+COPCO","Atlas Copco",s,flags=re.I)
    s=re.sub(r"(?<=\d),(?=\d)", ".", s)
    s=re.sub(r"\d{3,4}\s*/\s*[13]\s*/\s*\d{2}(\s*/\s*[A-ZА-Я]{1,3})?"," ",s)
    s=re.sub(r"\b\d{3,4}\s?[Вв]\b"," ",s)
    s=re.sub(r"\b[13]\s?ф\b"," ",s); s=re.sub(r"\b\d{2}\s?Гц\b"," ",s)
    s=re.sub(r"без\s*N\s*/?\s*CE"," ",s,flags=re.I)
    s=re.sub(r"\bN\s*/\s*CE\b|/?\s*\bCE\b"," ",s,flags=re.I)
    s=re.sub(r"\(\s*с\s*осушителем\s*\)"," ",s,flags=re.I)
    s=re.sub(r"[/,]"," ",s)
    s=re.sub(r"\b(220|230|380|400|660|690)\s+\d{1,2}(\s+YD)?\b"," ",s,flags=re.I)  # «400 50» (напряж+част), НЕ «500 10» (ресивер+бар)
    s=re.sub(r"\b(380|400|660|690)\b"," ",s); s=re.sub(r"\bYD\b"," ",s,flags=re.I)
    s=re.sub(r"(?<=\d)\.\s+(?=\d)",".",s)               # «10. 4» -> «10.4» (артефакт разбивки бара у ZR 55)
    s=re.sub(r"\b(bar|бар)\b"," ",s,flags=re.I)         # лишнее слово bar/бар внутри кода (бар добавим сами)
    s=re.sub(r"(?:^|(?<=\s))-(?=\s|$)"," ",s)           # одиночный дефис-токен «ZT 90 VSD - 10.4 FF»
    return re.sub(r"\s+"," ",s).strip(" -")
def kw_stem(s):                                          # основа для КЛЮЧА: без (...)-конфигов и спец-слов
    s=re.sub(r"\([^)]*\)"," ",s)                         # (на шасси)/(IP55)/(с осушителем)/(500)/(до -25С)
    s=re.sub(r"\bс\s+(осушител\w+|доохладител\w+|влагоотделит\w+|влагомаслоотделит\w+)\b"," ",s,flags=re.I)
    s=re.sub(r"\bна\s+раме\b|\bМоноблок\b|\bSkid\b|\bна\s+шасси\b|\bбез\s+шасси\b|\bв\s+кожухе\b"," ",s,flags=re.I)
    s=re.sub(r"\b\d+(?:[.,]\d+)?\s*к[Вв]т\b"," ",s,flags=re.I)   # «5 кВт» — мощность уже = название кампании
    return re.sub(r"\s+"," ",s).strip(" -")
def cap(s, n, bar=None):                                 # лимит Я.Директа 7 слов; бар (различитель давлений) не терять
    t=re.split(r"[ \-/]+", s.strip())                    # дефис/слэш у Директа = граница слова («5.5-10» -> 2)
    if len(t)<=n: return s                               # влезает — оставляем как есть (с дефисами)
    bt=("%g"%bar) if bar else None
    head=t[:n]
    if bt and (bt in t) and (bt not in head): head=t[:n-1]+[bt]   # выкидываем хвостовой спец-токен, бар оставляем
    return " ".join(head)

PRICE={}
for r in csv.DictReader(open(NP,encoding="utf-8-sig")):
    p=(r.get("price") or "").strip()
    if p: PRICE[nrm(r.get("product_url"))]=str(int(float(num(p))))
tr=list(csv.reader(open(TPL,encoding="utf-8-sig"),delimiter=";"))
HDR=tr[2]; TEMPLATE=tr[3]; NCOL=len(HDR)
FIXED={0,1,2,3,7,11,44,45,50,51,52,59}          # 10 (Название кампании) ставим сами
CENA=HDR.index("Цена")

prods=[]
for r in csv.DictReader(open(SRC,encoding="utf-8-sig"),delimiter=";"):
    nm=(r.get("Название") or "").strip()
    if not is_compressor(nm): continue
    bl=brand_from_text(nm) or find_brand(r.get("URL") or "")
    if bl=="enger": continue
    pw=num(r.get("Св-во: MOSHCHNOST_KVT") or "")
    nom=f"{nominal(pw):g} квт" if pw else "без мощности"
    bar=bar_value(r.get("Св-во: RABOCHEE_DAVLENIE_BAR"))
    core=clean_model(re.sub(r"\s+"," ",TYPE_RE.sub(" ",nm)).strip())
    url=(r.get("URL") or "").strip(); bp=r.get("Цена: Сайт (RUB)") or ""
    price=PRICE.get(nrm(url)) or (str(int(float(num(bp)))) if num(bp) else "")
    if not price: continue                           # без цены -> в кампанию не берём
    ipm=re.search(r"IP\s?(\d{2})", nm)
    e=(r.get("Св-во: CHASTOTNYY_PREOBRAZOVATEL") or "").strip().lower()=="да" or \
      bool(re.search(r"(?<![a-zа-яё])VSD(?![a-zа-яё])", nm, re.I))
    prods.append(dict(brand=disp(bl), core=core, bar=bar, url=url, price=price, nom=nom,
                      typ=product_type(nm), adj=type_adj(nm), ip=(ipm.group(1) if ipm else None), e=e))

def base_of(p):
    c=p["core"]
    if p["bar"]: c=re.sub(r"[-/ ]?"+re.escape(f"{p['bar']:g}")+r"\b","",c)
    c=re.sub(r"IP\s?\d{2}","",c,flags=re.I)
    c=re.sub(r"(?<![a-zа-яё])(VSD|частотник)(?![a-zа-яё])","",c,flags=re.I)
    return (p["brand"], re.sub(r"[ ,]+"," ",c).strip())
bbar=defaultdict(set); bip=defaultdict(set); be=defaultdict(set)
for p in prods:
    k=base_of(p)
    if p["bar"]: bbar[k].add(p["bar"])
    if p["ip"]: bip[k].add(p["ip"])
    be[k].add(p["e"])
def group_minus(p):
    k=base_of(p); parts=[]
    for b in sorted(bbar[k]-{p["bar"]}): parts.append(f"-{b:g}")
    for ip in sorted(bip[k]-({p["ip"]} if p["ip"] else set())): parts.append(f"-IP {ip}")
    if (not p["e"]) and (True in be[k]): parts += ["-частотник","-vsd","-инвертор"]
    return " ".join(parts)
def has_std_sibling(p): return False in be[base_of(p)]
def wc(s): return len(re.split(r"[ \-/]+", s.strip()))   # дефис/слэш — граница слова (как считает Я.Директ)

def build_phrases(p, broad, inner):
    # частотная версия при наличии стандартной — только по слову
    if p["e"] and has_std_sibling(p):
        out=[]; bw=cap(broad,6,p["bar"]); iw=cap(inner,6,p["bar"])   # +«частотник/vsd» -> остаёмся в 7 словах
        for w in ("частотник","vsd"):
            out.append((f"{bw} {w}",""))
            out.append((f"[{iw} {w}]",""))
        out.append(("---autotargeting","50")); return out
    adj=p["adj"]; tn="бустер" if p["typ"]=="бустер" else "компрессор"   # тип-существительное в ключ
    kompr=f"{tn} {broad}"; adjp=f"{adj} {tn} {broad}"   # «бустер X»/«компрессор X», «поршневой бустер X»
    kupit=f"{broad} купить"; cena=f"{broad} цена"
    hk=wc(kompr)<=7; ha=wc(adjp)<=7; hu=wc(kupit)<=7; hc=wc(cena)<=7
    bmin=[]                                     # кросс-минус на широкую: убрать пересечения
    if hk or ha: bmin.append(f"-{tn}")          # -компрессор/-бустер покрывает и «{adj} {tn}»
    if hu: bmin.append("-купить")
    if hc: bmin.append("-цена")
    out=[(broad+(" "+" ".join(bmin) if bmin else ""),""), (f"[{inner}]","")]
    if hk: out.append((kompr+(" -"+adj if ha else ""),""))
    if ha: out.append((adjp,""))
    if hu: out.append((kupit,""))
    if hc: out.append((cena,""))
    out.append(("---autotargeting","50")); return out

def make_row(p, nom, gnum, phrase, bid, h1, txt, disp_link, grp, gmin, first):
    row=[""]*NCOL
    for j in FIXED: row[j]=TEMPLATE[j]
    row[0]="-"; row[5]=grp; row[6]=str(gnum); row[10]=nom
    row[13]=phrase; row[15]=h1; row[16]="Компрессор Центр"; row[17]=txt
    row[40]=p["url"]; row[41]=disp_link; row[46]=bid; row[CENA]=p["price"]
    if first: row[60]=gmin
    return row

def build():
    os.makedirs(OUTDIR, exist_ok=True)
    for f in os.listdir(OUTDIR): os.remove(os.path.join(OUTDIR,f))
    camp=defaultdict(list)
    for p in prods: camp[p["nom"]].append(p)
    def sk(k): return (k=="без мощности", float(k.split()[0]) if k!="без мощности" else 0)
    made=[]; total=0; collapsed=0
    for nom in sorted(camp, key=sk):
        rows=[]; gnum=0; seen=set()                      # фразы кампании: дубли -> «конкуренция фраз», дедупим
        for p in camp[nom]:
            gnum+=1
            grp=p["core"] + (f" {p['bar']:g} бар" if p["bar"] and f"{p['bar']:g}" not in p["core"] else "")
            stem=kw_stem(p["core"])                      # основа ключа: без конфиг-хвостов
            broad=stem + (f" {p['bar']:g}" if p["bar"] and f"{p['bar']:g}" not in stem else "")
            broad=cap(broad, 7, p["bar"])                # держим лимит 7 слов, бар сохраняем
            inner=re.sub(r"\s+"," ",re.sub(r"[-/]"," ",broad)).strip()
            h1=headline(p["typ"], stem)                  # модель в заголовке — без (...)/конфигов
            txt=trimw(f"Надежный поставщик компрессоров {p['brand']} — нам доверяют лидеры рынка. Звоните!",81)
            disp_link=trimw(re.sub(r"\s+","-",broad),20); gmin=group_minus(p)
            kept=[]                                      # двойники (шасси/спец-суффикс) схлопываются в один ключ
            for ph,bid in build_phrases(p,broad,inner):
                if ph.startswith("---"): kept.append((ph,bid)); continue   # autotargeting — у каждой группы свой
                if ph in seen: continue                  # фраза уже есть в кампании -> не дублируем
                seen.add(ph); kept.append((ph,bid))
            if all(ph.startswith("---") for ph,_ in kept): collapsed+=1     # остались только на autotargeting
            for i,(ph,bid) in enumerate(kept):
                rows.append(make_row(p,nom,gnum,ph,bid,h1,txt,disp_link,grp,gmin,i==0))
        wb=openpyxl.Workbook(); ws=wb.active
        ws.append(["Предложение текстовых блоков для кампании"]+[""]*(NCOL-1)); ws.append(HDR)
        for r in rows: ws.append(r)
        for rr in range(3, ws.max_row+1):           # текст только там, где ведущий «-»
            ws.cell(rr,14).number_format="@"; ws.cell(rr,61).number_format="@"
        fn=os.path.join(OUTDIR, nom.replace(" ","_")+".xlsx"); wb.save(fn); made.append(fn)
        total+=len(rows)
    with zipfile.ZipFile(ZIP,"w",zipfile.ZIP_DEFLATED) as z:
        for p in made: z.write(p, os.path.basename(p))
    print(f"кампаний(файлов): {len(made)} | товаров: {len(prods)} | строк всего: {total}")
    print(f"групп-двойников (только autotargeting, ключ занят родственной моделью): {collapsed}")
    print(f"-> {ZIP}")

if __name__=="__main__":
    build()
