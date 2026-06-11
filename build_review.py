"""Собирает Excel для проверки: строка = товар prokompressor, в ячейках конкурентов
кликабельная цена (ссылка на их страницу).

Сопоставление по УМНОМУ отпечатку (matcher.signature): безопасные перестановки
склеиваются автоматически, структурные различия расходятся. Дедуп по URL -> мин. цена.
ЖЁЛТЫЙ = под один ключ у конкурента попало >1 разного URL (возможная неоднозначность)."""
import openpyxl, re, csv, sys
csv.field_size_limit(sys.maxsize)   # огромные specs-поля в прогонах парсера
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from matcher import signature, domain, find_brand
from scrape_files import SCRAPE_FILES
from atlas_need_specs import is_product_url   # отсев не-товарных (v-p-k /catalog/, аренда, б/у, статьи)

# Детектор «дробь в названии разная» (класс 5.5↔55, который URL не различает).
# Фильтруем Excel-битьё названий (1.46031 = сожранное 1,9/1,0) — там матч по URL верный.
_DEC=re.compile(r'\d+[.,]\d{1,2}(?!\d)')
_CORRUPT=re.compile(r'\d+[.,]\d{3,}')
def _clean_dec(name):
    return frozenset(m.group().replace(",",".")
                     for m in _DEC.finditer(re.sub(r'(?i)ip\s*\d+'," ",str(name))))
def _corrupt(name): return bool(_CORRUPT.search(str(name)))

SRC = "/root/.claude/uploads/62a19005-a7bf-569b-926e-b59b4a62600d/03d65e4a-_______________________.xlsx"
PROKO_CSV = "/root/.claude/uploads/62a19005-a7bf-569b-926e-b59b4a62600d/e7171060-products_export_20260608.csv"
SITEMAP = "/root/.claude/uploads/62a19005-a7bf-569b-926e-b59b4a62600d/63b1d773-all_sitemap_urls_1.xlsx"
# все прогоны парсера (прайс+статусы) — единый список в scrape_files.py (туда же добавляются новые)
CHECKED = SCRAPE_FILES
COMPETITORS = ["compressortyt.ru","aerocompressors.ru","pnevmoteh.ru",
               "pnevmo-sklad.ru","v-p-k.ru","rutector.ru"]
DOM_FIX = {"rostov.pnevmo-sklad.ru":"pnevmo-sklad.ru",
           "novosibirsk.pnevmo-sklad.ru":"pnevmo-sklad.ru"}
STATUS = {}   # url -> 'снято' / 'под заказ' (статус товара у конкурента)

def load_all():
    info={}; STATUS.clear()
    # прайс-файл: берём только КОНКУРЕНТОВ (наши товары — из свежего CSV ниже)
    ws = openpyxl.load_workbook(SRC, read_only=True, data_only=True)["Лист1"]
    for row in ws.iter_rows(min_row=1, values_only=True):
        u=str(row[12])
        if domain(u).replace("www.","")=="prokompressor.ru": continue
        name=str(row[0])
        try: p=float(str(row[3]).replace(",","."))
        except: p=None
        if u not in info: info[u]=[name,[]]
        if p is not None: info[u][1].append(p)
        if len(row)>8 and row[8] and "снят" in str(row[8]).lower(): STATUS[u]="снято"
    # свежий экспорт prokompressor (Название;Ссылка;Цена;Валюта) — наши товары и актуальные цены
    try:
        with open(PROKO_CSV, encoding="utf-8-sig", errors="replace") as fh:
            rd=csv.reader(fh, delimiter=";"); next(rd, None)
            for row in rd:
                if len(row)<3 or "prokompressor" not in row[1]: continue
                u=row[1].strip(); nm=row[0].strip().replace("&quot;",'"')
                try: p=float(str(row[2]).replace(",",".").replace(" ","")); p=p if p>0 else None
                except: p=None
                info[u]=[nm, ([p] if p else [])]
    except FileNotFoundError:
        pass
    # подмешиваем sitemap-URL конкурентов (без цены) — кандидаты на прайс-чек
    try:
        ws2 = openpyxl.load_workbook(SITEMAP, read_only=True, data_only=True)["Sheet1"]
        for row in ws2.iter_rows(min_row=2, values_only=True):
            u=str(row[0]); ul=u.lower()
            if u in info: continue
            if "kompressor" not in ul and "compressor" not in ul: continue
            d=domain(u).replace("www.",""); d=DOM_FIX.get(d,d)
            if d not in COMPETITORS: continue
            if not find_brand(u): continue
            seg=u.rstrip("/").split("/")[-1]
            if not any(ch.isdigit() for ch in seg): continue   # нет цифры в слаге = категория/серия, не товар
            info[u]=[seg, []]   # имя = слаг (цены нет)
    except FileNotFoundError:
        pass
    # присланные проверенные цены + статусы (series_status: снято/под заказ) — несколько прогонов
    for path in CHECKED:
        try:
            with open(path, encoding="utf-8-sig", errors="replace") as fh:
                for row in csv.DictReader(fh):
                    u=(row.get("product_url") or "").strip()
                    if not u: continue
                    p=None
                    try:                                     # ТОЛЬКО price: old_price=перечёркнутое
                        v=float(str(row.get("price","")).replace(",",".").replace(" ",""))
                        if v>0: p=v                           # «было» (устаревший мусор при «по запросу»)
                    except: pass
                    st=(row.get("series_status") or "").strip().lower()
                    if "снят" in st or st=="нет в наличии": STATUS[u]="снято"
                    elif st=="под заказ": STATUS.setdefault(u,"под заказ")
                    if p:
                        if u in info: info[u][1].append(p)
                        else: info[u]=[(row.get("name") or u).strip(), [p]]
        except FileNotFoundError:
            pass
    # отсев не-товарных страниц конкурентов (v-p-k /catalog/=категория со статусом «снято»
    # на уровне серии, аренда, б/у, статьи). Наш prokompressor /catalog/ — это товары, не трогаем.
    drop=[u for u in info if "prokompressor.ru" not in u and not is_product_url(u)]
    for u in drop: info.pop(u, None); STATUS.pop(u, None)
    return info

def cluster_all(info):
    """brand -> key -> domain -> {url: price}."""
    data=defaultdict(lambda: defaultdict(lambda: defaultdict(dict)))
    for u,(name,prices) in info.items():
        key=signature(u); brand=key.split("::")[0]
        d=domain(u); d=DOM_FIX.get(d,d)
        data[brand][key][d][u]=min(prices) if prices else None
    return data

HDR=["№","Отпечаток","Название","У нас","Конк-тов","Ваша цена","Ваша ссылка"]\
    +COMPETITORS+["min конк.","Δ к min, %","ВЕРДИКТ (ок / ошибка: ...)"]

def _pick_min(d):
    return min(d, key=lambda k:(d[k] is None, d[k]))

def _show_key(key):
    p=key.split("::")
    parts=[p[0]] + [x for x in p[1:] if x]
    return " · ".join(parts)

def _best_name(info, sites):
    """Лучшее имя среди конкурентов: длинное и без Excel-битья (.46xxx)."""
    cand=[u for c in COMPETITORS if c in sites for u in sites[c]]
    return info[max(cand, key=lambda u: len(info[u][0]) - (60 if ".46" in info[u][0] else 0))][0]

def add_sheet(wb, info, data, brand, sheet_name):
    bd=data.get(brand, {})
    wsx=wb.create_sheet(sheet_name[:31]); wsx.append(HDR)
    blue=Font(color="0563C1", underline="single")
    strike=Font(color="C00000", underline="single", strike=True)  # снято с производства
    bold=Font(bold=True, color="FFFFFF")
    hfill=PatternFill("solid", fgColor="305496")
    warn=PatternFill("solid", fgColor="FFE699")     # жёлтый: >1 разный URL у конкурента
    nomatch=PatternFill("solid", fgColor="F2F2F2")  # серый: не нашлось
    gapfill=PatternFill("solid", fgColor="DDEBF7")  # голубой: кандидат добавить (нас нет)
    checkfill=PatternFill("solid", fgColor="C6E0B4") # зелёный: есть у конкур., цены нет → проверить
    center=Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in range(1,len(HDR)+1):
        cell=wsx.cell(1,c); cell.font=bold; cell.fill=hfill; cell.alignment=center

    ncomp=lambda s: sum(1 for c in COMPETITORS if c in s)
    pk_keys=sorted(k for k,s in bd.items() if "prokompressor.ru" in s)
    gap_keys=sorted((k for k,s in bd.items()
                     if "prokompressor.ru" not in s and ncomp(s)>=2),
                    key=lambda k:(-ncomp(bd[k]), k))      # популярные — первыми
    n=0
    for key in pk_keys+gap_keys:
        sites=bd[key]; has_pk="prokompressor.ru" in sites; nc=ncomp(sites)
        if has_pk:
            pk=sites["prokompressor.ru"]; pk_url=_pick_min(pk)
            pk_price=pk[pk_url]; name=info[pk_url][0]
        else:
            pk_url=None; pk_price=None; name=_best_name(info, sites)
        adec=_clean_dec(name); acorr=_corrupt(name)
        n+=1; r=n+1
        wsx.cell(r,1,n)
        wsx.cell(r,2,_show_key(key))
        wsx.cell(r,3,name)
        uc=wsx.cell(r,4,"да" if has_pk else "нет")
        if not has_pk: uc.fill=gapfill
        cc=wsx.cell(r,5,nc)
        if not has_pk and nc>=4: cc.font=Font(bold=True)   # очень популярный кандидат
        if has_pk:
            c6=wsx.cell(r,6, pk_price if pk_price else "нет цены")
            if pk_price: c6.number_format="# ##0"
            l=wsx.cell(r,7,"открыть"); l.hyperlink=pk_url; l.font=blue
        else:
            wsx.cell(r,6,"—").fill=gapfill
        comp_prices=[]
        for ci,comp in enumerate(COMPETITORS):
            cell=wsx.cell(r,8+ci)
            urls=sites.get(comp)
            if urls:
                priced={u:p for u,p in urls.items() if p is not None}
                if priced:
                    cu=_pick_min(priced); pr=priced[cu]
                    cell.value=pr; cell.number_format="# ##0"; cell.hyperlink=cu
                    if STATUS.get(cu)=="снято":
                        cell.font=strike                      # снято — перечёркнуто, в min не считаем
                    else:
                        cell.font=blue; comp_prices.append(pr)
                        dec_bad=(has_pk and not acorr and not _corrupt(info[cu][0])
                                 and _clean_dec(info[cu][0])!=adec)   # риск занижения — только наши строки
                        if len(urls)>1 or dec_bad: cell.fill=warn   # неоднозначность / разная дробь
                else:                          # товар у конкурента есть, цены нет
                    cu=next(iter(urls))
                    if STATUS.get(cu)=="снято":
                        cell.value="снято"; cell.hyperlink=cu; cell.font=strike
                    else:
                        cell.value="проверить"; cell.hyperlink=cu; cell.font=blue; cell.fill=checkfill
            else:
                cell.fill=nomatch
        if comp_prices:
            mn=min(comp_prices)
            wsx.cell(r,14,mn).number_format="# ##0"
            if pk_price: wsx.cell(r,15, round((pk_price-mn)/mn*100,1))
    widths=[5,40,44,7,9,11,8]+[13]*len(COMPETITORS)+[11,10,24]
    for i,w in enumerate(widths,1): wsx.column_dimensions[get_column_letter(i)].width=w
    wsx.freeze_panes="D2"; wsx.auto_filter.ref=f"A1:{get_column_letter(len(HDR))}{n+1}"
    return n

def build_multi(brands, out_path):
    info=load_all(); data=cluster_all(info)
    wb=openpyxl.Workbook(); wb.remove(wb.active)
    for brand, sheet in brands:
        n=add_sheet(wb, info, data, brand, sheet); print(f"  {sheet:<14} {n} товаров")
    wb.save(out_path); print(f"-> {out_path}")

if __name__=="__main__":
    build_multi([("dalgakiran","Dalgakiran")], "/home/user/Statistic/Dalgakiran_review.xlsx")
