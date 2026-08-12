"""Четыре списка поверх матчинга, одной книгой xlsx.

    сводка          по брендам: сколько наших, сколько сматчено, сколько GAP
    снятые          позиции, которые конкуренты пометили «снято с производства»
    GAP 2+          чего у нас нет, а возят ДВА и более конкурента
    GAP 1           чего у нас нет, и возит только ОДИН — спрос не подтверждён

Почему GAP разделён по числу площадок. Одна площадка может завезти что угодно
под заказ; две и более — уже признак спроса. Проверка 11.08 подтверждает это
числом: среди позиций, которые держат 4-6 площадок, у нас отсутствует 11-18%,
а среди «только у двоих» — 59%. То есть чем шире позицию возят, тем вероятнее
она нужна и нам.

Одна позиция = один товар, а не одна карточка: карточки разных площадок
схлопываются той же логикой, что и матчинг (серия+номер, специи не противоречат,
исполнение сравнивается меткой).

    python tools/gap_lists.py                  # -> GAP_lists.xlsx
    python tools/gap_lists.py --out /tmp/g.xlsx
"""
from __future__ import annotations

import argparse
import os
import sys
from collections import defaultdict, Counter

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import brand_spec_review as B
import spec_match as S

HDR = PatternFill("solid", fgColor="DDEBF7")
WARN = PatternFill("solid", fgColor="FCE4D6")


def cname(c) -> str:
    return c.get("name") or B.slug(c["url"]).replace("_", " ")


def sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for i, c in enumerate(ws[1], 1):
        c.font = Font(bold=True); c.fill = HDR
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"
    return ws


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default="GAP_lists.xlsx")
    a = ap.parse_args()

    ours = B.load_ours_all(); cands = B.load_comp_all()

    # 1) что матчер уже связал — чтобы не назвать GAP то, что у нас есть
    matched_urls: set[str] = set()
    per_brand_matched: Counter = Counter()
    for brand in sorted(set(ours) & set(cands)):
        by = defaultdict(list)
        for c in cands[brand]:
            by[c["sn"]].append(c)
        for o in ours[brand]:
            m = B.pick_cands(o, by.get(o["sn"], []), brand)
            if m:
                per_brand_matched[brand] += 1
                matched_urls.update(c["url"] for c in m)

    # 2) схлопываем карточки конкурентов в товары
    products: list[tuple[str, list]] = []
    for brand, rows in cands.items():
        by_sn = defaultdict(list)
        for c in rows:
            by_sn[c["sn"]].append(c)
        for pool in by_sn.values():
            used = set()
            for i, x in enumerate(pool):
                if id(x) in used:
                    continue
                grp = [x]; used.add(id(x))
                xm = S.variant_letters(cname(x), brand)
                for y in pool[i + 1:]:
                    if id(y) in used or y["site"] == x["site"]:
                        continue
                    if not S.match(x, [y]):
                        continue
                    if not S.same_variant(xm, S.variant_letters(cname(y), brand)):
                        continue
                    grp.append(y); used.add(id(y))
                products.append((brand, grp))

    def row_of(brand, grp):
        sites = sorted({c["site"] for c in grp})
        pr = [c["price"] for c in grp if c.get("price")]
        best = min(grp, key=lambda c: c.get("price") or 10**12)
        return [brand, cname(best)[:180], len(sites), ", ".join(sites),
                min(pr) if pr else None, max(pr) if pr else None,
                best.get("kw"), best.get("bar"), best.get("fl"), best.get("rv"),
                best["url"]]

    gap2, gap1, sny = [], [], []
    per_brand_gap2, per_brand_gap1 = Counter(), Counter()
    for brand, grp in products:
        sites = {c["site"] for c in grp}
        snyato = [c for c in grp if (c.get("status") or "") == "снято"]
        if snyato:
            for c in snyato:
                sny.append([brand, cname(c)[:180], c["site"], c.get("price"),
                            "да" if any(u["url"] in matched_urls for u in grp) else "нет", c["url"]])
        if any(c["url"] in matched_urls for c in grp):
            continue                       # у нас есть — не GAP
        if len(sites) >= 2:
            gap2.append(row_of(brand, grp)); per_brand_gap2[brand] += 1
        else:
            gap1.append(row_of(brand, grp)); per_brand_gap1[brand] += 1

    gap2.sort(key=lambda r: (-r[2], r[0], r[1]))
    gap1.sort(key=lambda r: (r[0], r[1]))
    sny.sort(key=lambda r: (r[0], r[1]))

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    brands = sorted(set(ours) | set(cands))
    sheet(wb, "сводка",
          ["бренд", "наших карточек", "сматчено", "карточек у конкурентов",
           "GAP: 2+ площадки", "GAP: 1 площадка", "снятых у конкурентов"],
          [[b, len(ours.get(b, [])), per_brand_matched[b], len(cands.get(b, [])),
            per_brand_gap2[b], per_brand_gap1[b],
            sum(1 for r in sny if r[0] == b)] for b in brands],
          [18, 15, 12, 20, 16, 16, 18])

    cols = ["бренд", "товар", "площадок", "где есть", "цена мин", "цена макс",
            "кВт", "бар", "л/мин", "ресивер", "ссылка"]
    w = [16, 60, 10, 34, 13, 13, 8, 7, 9, 9, 60]
    sheet(wb, "GAP 2+ площадки", cols, gap2, w)
    sheet(wb, "GAP 1 площадка", cols, gap1, w)
    sheet(wb, "снятые у конкурентов",
          ["бренд", "товар", "площадка", "цена", "есть у нас", "ссылка"], sny,
          [16, 70, 22, 13, 12, 60])

    wb.save(a.out)
    print(f"брендов: {len(brands)}")
    print(f"товаров у конкурентов (после схлопывания): {len(products):,}")
    print(f"  GAP, возят 2+ площадки: {len(gap2):,}")
    print(f"  GAP, возит 1 площадка:  {len(gap1):,}")
    print(f"  снятых с производства:  {len(sny):,}")
    print(f"\n-> {a.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
