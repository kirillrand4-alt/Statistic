"""Оставить в присланном файле только товары, у которых цены на сайте НЕТ.

    python filtr.py <исходный.xlsx> <лист> <номер колонки с именем> <результат обхода.csv> <итог.xlsx>

«Цены на сайте нет» — это вердикты «цена по запросу», «404» и «нет в каталоге».
Строки с вердиктом «цена в блоке товара» удаляются: там пометка «нет цены» неверна.

Почему не `delete_rows`: у заказчика на каждой ячейке висит гиперссылка на карточку
(3 007 штук во втором файле), а `delete_rows` двигает значения и стили, но НЕ двигает
гиперссылки. При сохранении openpyxl привязывает осиротевшую ссылку к пустой ячейке и
подставляет её адрес значением — в книге появлялись 764 строки с голыми URL. Поэтому
снимаем нужные строки целиком (значение + стиль + ссылка), чистим лист и кладём обратно.
"""
import csv
import sys
from collections import Counter

import openpyxl

XLSX, SHEET, NCOL, REZ, OUT = sys.argv[1], sys.argv[2], int(sys.argv[3]), sys.argv[4], sys.argv[5]

verdict = {}
for r in csv.DictReader(open(REZ, encoding="utf-8-sig"), delimiter=";"):
    verdict[r["товар"]] = r["вердикт"]
with_price = {k for k, v in verdict.items() if v == "цена в блоке товара"}

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]
was = ws.max_row
keep = [i for i in range(2, was + 1)
        if str(ws.cell(i, NCOL + 1).value or "") not in with_price]

# снимок до чистки: стиль берём объектом _style, ссылку — объектом Hyperlink
snap = [[(c.value, c._style, c.hyperlink) for c in ws[i]] for i in keep]
names = [str(ws.cell(i, NCOL + 1).value or "") for i in keep]

ws.delete_rows(2, was)
for n, row in enumerate(snap, start=2):
    for j, (val, style, link) in enumerate(row, start=1):
        c = ws.cell(n, j)
        c.value = val
        c._style = style
        if link is not None:
            link.ref = c.coordinate      # ref прибит к старой строке, переставляем
            c._hyperlink = link

# вердикт обхода последней колонкой: «404» — это мёртвая карточка, а не «цена по запросу»,
# и лечится это по-разному, поэтому их надо различать глазами.
vc = ws.max_column + 1
ws.cell(1, vc, "Проверка сайта").font = ws.cell(1, 1).font.copy()
for n, name in enumerate(names, start=2):
    ws.cell(n, vc, verdict.get(name, "не проверено"))

if ws.auto_filter.ref:
    ws.auto_filter.ref = f"A1:{ws.cell(1, vc).column_letter}{ws.max_row}"
wb.save(OUT)

print(f"исходных строк: {was - 1}")
print(f"удалено (цена на сайте есть): {was - 1 - len(keep)}")
print(f"осталось (цены на сайте нет): {len(keep)}")
print("вердикты оставшихся:", dict(Counter(verdict.get(n, "не проверено") for n in names)))
print("->", OUT)
