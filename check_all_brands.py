"""Сводный лист «Проверить карточку» по ВСЕМ брендам: наши товары, где характеристика
расходится с консенсусом ≥2 сайтов конкурентов (та же модель: серия+номер, остальные
поля запиннены). Сортировка по КРАТНОСТИ расхождения — ×N сверху (Atmos производ. ×10).
Считаем только по НЕсматченным (у сматченных спека уже сошлась)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from collections import defaultdict
from brand_spec_review import load_ours_all, load_comp_all, FIELDS
from spec_match import match, receiver_filter

OUT="/home/user/Statistic/Proverit_kartochku_ALL.xlsx"

def find_issue(o, same):
    for label,key,tol in FIELDS:
        ov=o.get(key)
        if not ov: continue
        others=[(k2,t2) for (l2,k2,t2) in FIELDS if k2!=key]
        variant=[c for c in same if c.get(key)
                 and (c["ff"] or 0)==(o.get("ff") or 0) and (c["vsd"] or 0)==(o.get("vsd") or 0)
                 and all(o.get(k2) and c.get(k2) and abs(o[k2]-c[k2])<=t2*max(o[k2],c[k2])
                         for k2,t2 in others)]
        for c1 in variant:
            v1=c1[key]; doms={c2["site"] for c2 in variant if abs(c2[key]-v1)<=tol*max(c2[key],v1)}
            if len(doms)>=2 and abs(ov-v1)>tol*max(ov,v1):
                src=next(c2 for c2 in variant if abs(c2[key]-v1)<=tol*max(c2[key],v1))
                ratio=max(ov,v1)/min(ov,v1)
                return (label, ov, v1, len(doms), ratio, src)
    return None

def build():
    ours_all=load_ours_all(); cands_all=load_comp_all()
    rows=[]
    for b in set(ours_all)&set(cands_all):
        by=defaultdict(list)
        for c in cands_all[b]: by[c["sn"]].append(c)
        for o in ours_all[b]:
            if receiver_filter(o.get("rv"), match(o, by.get(o["sn"], []))): continue  # сматчился
            iss=find_issue(o, by.get(o["sn"], []))
            if iss: rows.append((b, o, iss))
    rows.sort(key=lambda t:-t[2][4])   # по кратности убыв.

    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Проверить карточку"
    blue=Font(color="0563C1", underline="single"); orange=Font(color="C55A11", underline="single")
    red=Font(color="C00000", bold=True); bold=Font(bold=True, color="FFFFFF")
    hfill=PatternFill("solid", fgColor="305496"); redfill=PatternFill("solid", fgColor="FFC7CE")
    center=Alignment(horizontal="center", vertical="center", wrap_text=True)
    HDR=["№","Бренд","Наш товар","Ваша цена","Поле","У нас","У конкур.","Крат.","Сайтов","Подтверждение (конкурент)"]
    ws.append(HDR)
    for ci in range(1,len(HDR)+1):
        c=ws.cell(1,ci); c.font=bold; c.fill=hfill; c.alignment=center
    r=1
    for b,o,(label,ov,v1,nd,ratio,src) in rows:
        r+=1
        ws.cell(r,1,r-1); ws.cell(r,2,b)
        ws.cell(r,3,o["name"])
        c4=ws.cell(r,4, o["price"] if o["price"] else "нет цены")
        if o["price"]: c4.number_format="# ##0"
        c4.hyperlink=o["url"]; c4.font=blue
        ws.cell(r,5,label)
        ws.cell(r,6, round(ov,1)); ws.cell(r,7, round(v1,1))
        rc=ws.cell(r,8, f"×{ratio:.1f}")
        if ratio>=2: rc.font=red; rc.fill=redfill          # расхождение в РАЗЫ — критично
        ws.cell(r,9, nd)
        lk=ws.cell(r,10, f"[{src['site']}] {src['name'][:48]}"); lk.hyperlink=src["url"]; lk.font=orange
    widths=[5,13,46,11,8,9,10,7,7,52]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="C2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(HDR))}{r}"
    wb.save(OUT)
    from collections import Counter
    print(f"всего: {len(rows)} | в разы (×2+): {sum(1 for _,_,i in rows if i[4]>=2)}")
    print("по полям:", dict(Counter(i[0] for _,_,i in rows)))
    print("топ брендов:", Counter(b for b,_,_ in rows).most_common(8))
    print(f"-> {OUT}")

if __name__=="__main__":
    build()
