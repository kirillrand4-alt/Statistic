"""Бренды, которых НЕТ у нас, но есть на 2+ сайтах конкурентов — каталог их моделей
с ценами по сайтам (кандидаты «завести»). По xlsx на бренд -> архив.
Группировка по физ-ключу (исполнения AC/WC/Pack сливаются, цена минимальная)."""
import os, zipfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from collections import defaultdict
from brand_spec_review import load_comp_all, COMPETITORS

TARGETS = ["gmp","alup","xeleron","denair","brestor","vortex","kaishan","atom","baldor","compair"]
OUTDIR  = "/home/user/Statistic/gap_reports"
ZIP     = "/home/user/Statistic/Gap_brands_catalog.zip"

def build():
    os.makedirs(OUTDIR, exist_ok=True)
    cands_all=load_comp_all()
    blue=Font(color="0563C1", underline="single"); strike=Font(color="C00000", underline="single", strike=True)
    bold=Font(bold=True, color="FFFFFF"); hfill=PatternFill("solid", fgColor="305496")
    warn=PatternFill("solid", fgColor="FFE699"); nomatch=PatternFill("solid", fgColor="F2F2F2")
    center=Alignment(horizontal="center", vertical="center", wrap_text=True)
    made=[]
    for b in TARGETS:
        cands=cands_all.get(b, [])
        if not cands: print(f"  {b:<12} — нет кандидатов, пропуск"); continue
        groups=defaultdict(lambda: defaultdict(list))      # gkey -> site -> [cards]
        for c in cands:
            gk=(c["sn"], round(c["kw"]) if c["kw"] else None, round(c["bar"]) if c["bar"] else None,
                c["ff"] or 0, c["vsd"] or 0, c["rv"])
            groups[gk][c["site"]].append(c)
        rows=sorted(groups.items(), key=lambda t:(-len(t[1]),
                    -sum(1 for _ in t[1].values())))       # больше сайтов — выше
        wb=openpyxl.Workbook(); ws=wb.active; ws.title="каталог"
        HDR=["№","Модель","Серия","кВт","бар","произв","ресивер","FF","VSD","Сайтов"]+COMPETITORS+["min цена"]
        ws.append(HDR)
        for ci in range(1,len(HDR)+1):
            cc=ws.cell(1,ci); cc.font=bold; cc.fill=hfill; cc.alignment=center
        r=1
        for gk,sites in rows:
            r+=1; allc=[c for cs in sites.values() for c in cs]
            ws.cell(r,1,r-1); ws.cell(r,2, max(allc,key=lambda c:len(c["name"]))["name"][:70])
            ws.cell(r,3, f"{str(gk[0][0]).upper()}{gk[0][1]:g}")
            ws.cell(r,4,gk[1] or ""); ws.cell(r,5,gk[2] or ""); ws.cell(r,6, allc[0]["fl"] or "")
            ws.cell(r,7, gk[5] if gk[5] else ""); ws.cell(r,8,"FF" if gk[3] else ""); ws.cell(r,9,"VSD" if gk[4] else "")
            ws.cell(r,10,len(sites)); prices=[]
            for ci,site in enumerate(COMPETITORS):
                cell=ws.cell(r,11+ci); cs=sites.get(site)
                if not cs: cell.fill=nomatch; continue
                priced=[c for c in cs if c["price"] and c["status"]!="снято"]
                show=min(priced,key=lambda c:c["price"]) if priced else cs[0]
                if show["price"]:
                    cell.value=show["price"]; cell.number_format="# ##0"; cell.hyperlink=show["url"]
                    cell.font=strike if show["status"]=="снято" else blue
                    if show["status"]!="снято": prices.append(show["price"])
                else:
                    cell.value="снято" if show["status"]=="снято" else "По запросу"
                    cell.hyperlink=show["url"]; cell.font=strike if show["status"]=="снято" else blue
                if len(cs)>1: cell.fill=warn
            if prices: ws.cell(r,17,min(prices)).number_format="# ##0"
        widths=[5,50,10,7,7,9,9,5,5,8]+[13]*6+[11]
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
        ws.freeze_panes="C2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(HDR))}{r}"
        path=os.path.join(OUTDIR, f"{b.upper()}_katalog.xlsx"); wb.save(path); made.append(path)
        print(f"  {b:<12} моделей {r-1:>4}")
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in made: z.write(p, os.path.basename(p))
    print(f"-> {ZIP} ({len(made)} брендов)")

if __name__=="__main__":
    build()
