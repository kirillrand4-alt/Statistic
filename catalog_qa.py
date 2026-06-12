"""Внутренняя QA нашего каталога: число в НАЗВАНИИ товара противоречит нашему же СПЕК-полю.
Находка аудита v4 (подтверждена). Не баг матчинга (матчер берёт имя), а ошибка данных
в Битриксе — спек-поле заполнено неверно. Классы: ресивер (имя≠объём-спека), давление
(имя≠бар-спека). Фильтр ложных: компрессоры высокого давления (Paramina 350бар, имя=л.с.)."""
import csv, sys, re
csv.field_size_limit(sys.maxsize)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from spec_match import text_flags, num, sane_bar, bar_value, is_compressor
from scrape_files import U

SPECS_CSV=U+"specs2/specs_compact.csv"; PROKO_CSV=U+"e7171060-products_export_20260608.csv"
OUT="/home/user/Statistic/Catalog_internal_qa.xlsx"

def build():
    rows={}
    for r in csv.DictReader(open(SPECS_CSV, encoding="utf-8-sig", errors="replace"), delimiter=";"):
        code=(r.get("IE_CODE") or "").strip()
        if not code: continue
        cur=rows.setdefault(code,{})
        for k,v in r.items():
            if v and not cur.get(k): cur[k]=v.strip()
    price={}
    for row in csv.reader(open(PROKO_CSV, encoding="utf-8-sig", errors="replace"), delimiter=";"):
        if len(row)<3 or "prokompressor" not in row[1]: continue
        sl=row[1].rstrip("/").split("/")[-1].lower()
        price[sl]=row[1].strip()
    rec=[]; pre=[]
    for code,r in rows.items():
        name=r.get("IE_NAME","")
        if not is_compressor(name+" "+code): continue
        url=price.get(code.lower(), f"https://prokompressor.ru/catalog/{code}/")
        # ресивер: имя vs спека «Объём ресивера, л»
        nrv=text_flags(name+" "+code)[2]
        srv=num(r.get("IP_PROP22564"))
        if nrv and nrv>1 and srv and abs(nrv-srv)>0.05*max(nrv,srv):
            rec.append((name, nrv, srv, url))
        # давление: имя «13FF/10P/14бар» vs спека бар
        m=re.search(r'[-/ ](\d{1,2})\s?(?:ff|p|бар|bar)\b', name.lower().replace(",","."))
        nbar=sane_bar(float(m.group(1))) if m else None
        sbar=bar_value(r.get("IP_PROP22573"))
        # ложные: компрессоры ВД (спека>40 бар, имя<40 = л.с./ступени) — пропускаем
        if nbar and sbar and abs(nbar-sbar)>0.05*max(nbar,sbar) and not (sbar>40 and nbar<40):
            pre.append((name, nbar, sbar, url))

    wb=openpyxl.Workbook(); wb.remove(wb.active)
    blue=Font(color="0563C1", underline="single"); bold=Font(bold=True,color="FFFFFF")
    red=Font(color="C00000", bold=True); hf=PatternFill("solid",fgColor="305496")
    ce=Alignment(horizontal="center",wrap_text=True)
    for title, data, lab in (("Ресивер имя≠спека", rec, "Объём ресивера, л"),
                             ("Давление имя≠спека", pre, "Давление, бар")):
        ws=wb.create_sheet(title)
        HDR=["№","Наш товар","В названии","В спек-поле ("+lab+")","Расхождение"]
        ws.append(HDR)
        for i in range(1,len(HDR)+1):
            c=ws.cell(1,i); c.font=bold; c.fill=hf; c.alignment=ce
        for n,(name,a,b,url) in enumerate(sorted(data, key=lambda x:-max(x[1],x[2])/min(x[1],x[2])),1):
            r=n+1
            ws.cell(r,1,n)
            c=ws.cell(r,2,name); c.hyperlink=url; c.font=blue
            ws.cell(r,3,a); ws.cell(r,4,b)
            rc=ws.cell(r,5, f"×{max(a,b)/min(a,b):.2f}"); rc.font=red
        for i,w in enumerate([5,56,12,24,12],1): ws.column_dimensions[get_column_letter(i)].width=w
        ws.freeze_panes="C2"; ws.auto_filter.ref=f"A1:E{len(data)+1}"
    wb.save(OUT)
    print(f"ресивер имя≠спека: {len(rec)} | давление имя≠спека (без ВД-ложных): {len(pre)}")
    print(f"-> {OUT}")

if __name__=="__main__":
    build()
