"""Atlas-компрессоры конкурентов, у которых характеристики НЕ сняты или сняты НЕПОЛНО
(старый парсер). «Полно» = есть кВт И бар И производительность (три ключа спек-сцепки).

Вселенная (все Atlas-компрессоры конкурентов): прайс-файл + sitemap + ВСЕ прогоны парсера.
Полнота: specs объединяем по URL из всех прогонов — если характеристику сняли хоть в одном
прогоне, считаем её снятой (поздний прогон переписывает ключ). Итог -> txt для дочистки."""
import csv, sys, json, re
csv.field_size_limit(sys.maxsize)
import openpyxl
from collections import Counter
from urllib.parse import urlparse, unquote
from matcher import brand_of, domain
from spec_match import is_compressor
from scrape_files import SCRAPE_FILES, NEW_PARSER_FILES, U

SRC     = U + "03d65e4a-_______________________.xlsx"          # большой прайс-файл конкурентов
SITEMAP = U + "63b1d773-all_sitemap_urls_1.xlsx"
OUT     = "/home/user/Statistic/atlas_competitors_need_specs.txt"
COMPETITORS = {"compressortyt.ru","aerocompressors.ru","pnevmoteh.ru",
               "pnevmo-sklad.ru","v-p-k.ru","rutector.ru"}
DOM_FIX = {"rostov.pnevmo-sklad.ru":"pnevmo-sklad.ru","novosibirsk.pnevmo-sklad.ru":"pnevmo-sklad.ru"}
def dm(u): d=domain(u).replace("www.",""); return DOM_FIX.get(d,d)
def slug(u):  # последний сегмент пути (без домена: 'aerocompressors' не должен ложно давать «compressor»)
    p=unquote(urlparse(str(u)).path).rstrip("/")
    return p.split("/")[-1] if p else ""

def canon_key(u):
    """Ключ одного товара: www/региональные субдомены (novosibirsk./rostov.) и хвостовой
    слеш не различаем — одна страница на разных хостах = один товар."""
    p=urlparse(str(u).lower())
    host=re.sub(r"^(www|novosibirsk|rostov)\.","",p.netloc)
    return host+p.path.rstrip("/")

def dedup_urls(urls):
    """Уникальные товары по canon_key; из вариантов берём короткий (обычно www/основной)."""
    best={}
    for u in urls:
        k=canon_key(u)
        if k not in best or len(u)<len(best[k]): best[k]=u
    return sorted(best.values())

def is_product_url(u):
    """Отсев НЕ-товарных страниц (аренда / категории-листинги / статьи / проекты / б-у),
    которые проходят по бренду+числу."""
    ul=str(u).lower()
    if "arenda" in ul: return False             # аренда оборудования, не продажа
    if "v-p-k.ru/catalog" in ul: return False   # v-p-k: /catalog/=категория; товары в /product/
    if re.search(r"/(stati|articles|blog|news|novosti|company|nashi_proekty|proekty)/", ul):
        return False                            # статьи-обзоры / новости / кейсы «наши проекты»
    if "bu-oborud" in ul or re.search(r"(?:^|[-_/])b[-_]?u(?:[-_/0-9]|$)|bez_narabotki", ul):
        return False                            # б/у: цена несопоставима с новым
    return True

# --- наличие характеристики в specs (СНЯЛ ли парсер число; не матчинг-санити!) ---
# Для «надо ли дочистить» важно, есть ли число вообще. Диапазоны/большие промышленные
# значения (ZR900: '56,5 - 142,7') = снято; матчинг-санити (кап 120000 л/мин) тут не к месту.
def _hasnum(v): return bool(re.search(r"\d", str(v)))
def has_kw(d):
    return any("мощ" in k.lower() and "шум" not in k.lower() and "звук" not in k.lower() and _hasnum(v)
              for k,v in d.items())
def has_bar(d):  return any("давлен" in k.lower() and _hasnum(v) for k,v in d.items())
def has_flow(d):
    from spec_match import is_flow_key   # «Производитель»/«Производство» ≠ производительность
    return any(is_flow_key(k) and _hasnum(v) for k,v in d.items())

def best_name(cur, nm):
    """Длинное имя без Excel-битья (.46xxx) — приоритет."""
    if not nm: return cur
    if not cur: return nm
    return nm if (len(nm) > len(cur) and ".46" not in nm) else cur

def load_universe():
    """Вся вселенная URL конкурентов: names (лучшее имя), specs (объединённые по URL,
    поздний прогон переписывает ключи), rescanned (сканированы ОБНОВЛЁННЫМ парсером)."""
    names={}; specs={}; rescanned=set()
    # 1) specs + имена из всех прогонов парсера
    for f in SCRAPE_FILES:
        try: fh=open(f, encoding="utf-8-sig", errors="replace")
        except FileNotFoundError: continue
        new_era = f in NEW_PARSER_FILES
        for r in csv.DictReader(fh):
            u=(r.get("product_url") or "").strip()
            if not u: continue
            if new_era: rescanned.add(u)
            names[u]=best_name(names.get(u,""), (r.get("name") or "").strip())
            sp=r.get("specs") or ""
            if sp:
                try: d=json.loads(sp)
                except Exception: d=None
                if isinstance(d, dict) and d: specs.setdefault(u,{}).update(d)
    # 2) URL-вселенная из прайс-файла и sitemap (имена; specs там нет)
    try:
        ws=openpyxl.load_workbook(SRC, read_only=True, data_only=True)["Лист1"]
        for row in ws.iter_rows(min_row=1, values_only=True):
            u=str(row[12]); names[u]=best_name(names.get(u,""), str(row[0]))
    except (FileNotFoundError, KeyError): pass
    try:
        ws2=openpyxl.load_workbook(SITEMAP, read_only=True, data_only=True)["Sheet1"]
        for row in ws2.iter_rows(min_row=2, values_only=True):
            u=str(row[0]); names.setdefault(u, u.rstrip("/").split("/")[-1])
    except (FileNotFoundError, KeyError): pass
    return names, specs, rescanned

def build():
    names, specs, rescanned = load_universe()
    # 3) фильтр Atlas-компрессоров конкурентов + проверка полноты характеристик.
    # «На странице нет данных»: URL пересканирован ОБНОВЛЁННЫМ парсером, поле всё равно
    # пустое -> сайт его не публикует, в txt не включаем (для матчинга пусто=не противоречит).
    need=[]; nodata=[]; complete=0
    by_site=Counter(); reason=Counter()
    for u,nm in names.items():
        if dm(u) not in COMPETITORS: continue
        if not is_product_url(u): continue
        if brand_of(u, nm) != "atlas": continue
        text=(nm or "")+" "+slug(u)               # имя+слаг, БЕЗ домена
        if not any(ch.isdigit() for ch in text): continue   # нет числа = категория/серия
        if not is_compressor(text): continue
        sd=specs.get(u, {})
        kw=has_kw(sd); bar=has_bar(sd); fl=has_flow(sd)
        if kw and bar and fl:
            complete+=1; continue
        if u in rescanned:
            nodata.append(u); continue
        need.append(u); by_site[dm(u)]+=1
        if not sd:                 reason["нет данных вообще (не сканировали)"]+=1
        else:
            if not kw:  reason["нет мощности кВт"]+=1
            if not bar: reason["нет давления бар"]+=1
            if not fl:  reason["нет производительности"]+=1

    need=sorted(set(need))
    with open(OUT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(need)+"\n")
    print(f"Atlas-компрессоры конкурентов: всего {complete+len(need)+len(nodata)} | "
          f"полные {complete} | сайт не публикует {len(nodata)} | НАДО ДОЧИСТИТЬ {len(need)}")
    print("\nпо сайтам (надо дочистить):")
    for s,c in by_site.most_common(): print(f"  {s:<20} {c}")
    print("\nпричины:")
    for s,c in reason.most_common(): print(f"  {s:<35} {c}")
    print("\nсайт не публикует поле (пересканировано новым парсером, в txt НЕ включены):")
    for u in sorted(nodata): print("  ", u)
    print(f"\n-> {OUT}")

if __name__=="__main__":
    build()
