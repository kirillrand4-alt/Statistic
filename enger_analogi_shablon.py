"""Аналоги Enger по ШАБЛОНУ СВОЙСТВ заказчика — отчёт в раскладке присланного образца.

Отличие от `enger_analogi.py` (он остаётся как был):

  * конкуренты берутся НЕ с площадок-перекупщиков, а с сайтов производителей
    (`vendor_data/*.csv`, собрано `tools/vendor_scrape.py`), и с каждого сайта — только
    его бренд. Это прямое требование заказчика: цена нужна заводская;
  * сцепка не по нашим допускам (±3% давление, ±5% производительность), а по 13
    свойствам шаблона с ТОЧНЫМ совпадением — допусков нет;
  * десять «зелёных» свойств шаблона при молчании сайта заменяются первым значением
    столбца, как и просил заказчик. Разбор цветов и обоснование — в `tools/enger_specs.py`.

Листа со сводкой два, и это не дубль:

  * «Сводная» — раскладка, формулы и цвета образца один в один: строка = модель Enger,
    две строки на модель (во второй следующий по цене аналог бренда), сцепка по 8 бар с
    откатом по ряду, если такого давления у серии нет (LUF идёт на 7,5 / 8,6 / 10,4);
  * «Сводная по барам» — строка на каждое давление. Нужна потому, что цена у нас по
    барам различается: BS-5,5BFR-250 на 8 бар стоит 312 125 ₽, на остальных 327 844 ₽,
    и модельная строка образца показывает одну цену из пяти.

    python enger_analogi_shablon.py            # боевой прогон
    python enger_analogi_shablon.py --proba    # только 10 моделей из образца, с разбором

Данные: OURS_DIR (наш каталог), vendor_data/ (сайты производителей).
"""
import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

_HERE = Path(__file__).resolve().parent          # скрипт зовут и из другой рабочей папки
sys.path[:0] = [str(_HERE), str(_HERE / "tools")]

import openpyxl
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter

from brand_spec_review import load_ours_all
from scrape_files import find_ours
from enger_specs import (PROPS, DEFAULTS, KW_LADDER, ip_classes, bar_bucket, bar_classes,
                         kw_value, num,
                         pick, yesno, yesno_or_model, from_name, block_class, motor_class,
                         cooling, stages)

VENDORS = Path("/home/user/Statistic/vendor_data")
OUT = "/home/user/Statistic/Enger_analogi_shablon.xlsx"

# Порядок брендов — как в образце («Сводная», тройки колонок с AA). Remeza оставлен на
# своём месте пустым: remeza.com — сайт завода, цен там нет ни в разделах, ни в карточках
# (проверено 24.08 обходом /catalog/compressors/screw_compressors/), а брать её цену с
# площадки нельзя — заказчик просил заводские сайты.
BRANDS = ["XELERON", "HANSMANN", "MAGNUS", "ET", "GMP", "CROSSAIR", "BERG", "ATOM",
          "KRAFTMACHINE", "EXELUTE", "REMEZA", "IRONMAC",
          # ниже — бренды, которых в образце нет, но сайты собраны и заказчик их называл
          "DALI", "COMPRAG", "SOLLANT", "ЗИФ"]
TITLE = {"XELERON": "Xeleron", "HANSMANN": "Hansmann", "MAGNUS": "MAGNUS", "ET": "ET",
         "GMP": "GMP", "CROSSAIR": "Cross-air", "BERG": "BERG", "ATOM": "ATOM",
         "KRAFTMACHINE": "Kraftmachine", "EXELUTE": "Exelute", "REMEZA": "Remeza",
         "IRONMAC": "Ironmac", "DALI": "Dali", "COMPRAG": "Comprag",
         "SOLLANT": "Sollant", "ЗИФ": "ЗИФ"}

# Модели из листа «Сводная» образца — контрольная выборка: на ней сверяем, какой режим
# подстановки умолчаний воспроизводит ручную сцепку заказчика.
PROBA = ["BS-5,5BFR-250", "HC-7,5BFR-250", "HC-7,5DFRE-250", "HC-11BFR-400",
         "HC-11DFRE-400", "HC-15BFR-400", "HC-15DFRE-400", "HC-18,5BFR-450",
         "HC-22BFR-450", "HC-22DFRE-450"]


# --- наша сторона -----------------------------------------------------------------------
def our_ip_by_name() -> dict:
    """IP из выгрузки Битрикса (22959) по имени карточки.

    load_ours_all() это свойство в матчер НЕ пускает — арбитраж 12.08 дал 15 ложных
    отсевов из 18 (PROPS_BITRIX.md). Но в шаблоне заказчика степень защиты — колонка
    ключа, и без неё отчёт не построить. Компромисс: значение показываем и используем
    как тристейт (режем только при явном конфликте обеих сторон), а расхождения выносим
    отдельной пометкой в строке."""
    path = find_ours("specs_compact")
    out = {}
    for r in csv.DictReader(open(path, encoding="utf-8-sig", errors="replace"), delimiter=";"):
        name = (r.get("IE_NAME") or "").strip()
        if not name:
            continue
        raw = (r.get("IP_PROP22959") or "").strip() or (r.get("IP_PROP22674") or "").strip()
        cls = ip_classes(raw)
        if cls:
            out[name] = cls
    return out


_ENG_TAIL = re.compile(r"\s+(\d+(?:[.,]\d+)?)\s*$")


def enger_family(name: str, bar) -> str:
    """«Винтовой компрессор Enger BS-5,5BFR-250 8» → «BS-5,5BFR-250».

    Хвостовое число отрезаем ТОЛЬКО если оно совпало с давлением карточки: у части серий
    хвост — часть кода («BS-11DFR-400-5»), и слепое отрезание сломало бы имя модели."""
    s = re.sub(r"^.*?\bEnger\b\s*", "", name, flags=re.I).strip()
    m = _ENG_TAIL.search(s)
    if m and bar and abs(float(m.group(1).replace(",", ".")) - float(bar)) < 0.05:
        s = s[:m.start()].strip()
    return s


def enger_props(o: dict, ip_map: dict) -> dict:
    """13 свойств шаблона для нашей карточки. None = «в выгрузке не указано»."""
    p = dict.fromkeys(PROPS)
    nm = o.get("name") or ""
    p["ресивер"] = "да" if o.get("rv") else None
    p["осушитель"] = "да" if o.get("ff") == 1 else None
    p["частотник"] = {1: "да", 0: "нет"}.get(o.get("vsd"))
    oil = o.get("oil")
    p["безмасляный"] = {"безмасло": "да", "масл": "нет"}.get(oil)
    p["охлаждение"] = {"возд": "воздушный", "вода": "водяной"}.get(o.get("cool"))
    p.update(from_name(nm))          # ступени / передвижной / PM-двигатель / безмасляный из имени
    # «блок» и «фильтры» в выгрузке Битрикса отсутствуют как свойства (22670 и 22760 в
    # компакт не выгружаются) — по всему нашему каталогу это честное «не указано».
    return {"kw": kw_value(o.get("kw")), "bar": bar_bucket(o.get("bar")),
            "ip": ip_map.get(nm), **p}


# --- сторона сайтов производителей ------------------------------------------------------
# Хвосты кодов моделей, которые сайт не дублирует в характеристиках. Каждый выведен из
# самих данных, а не из общих соображений (правило 3): рядом на сайте лежат обе версии.
_BERG_RCV = re.compile(r"-(\d{2,4})\s*$")                 # ВК-5.5Р-500 → ресивер 500 л
_BERG_DRY = re.compile(r"[А-Яа-я]*О-\d{2,4}\s*$")         # ВК-5.5РО-500 → «О» = осушитель
_KM_DRY = re.compile(r"/\s*О\b", re.I)                    # КМ5.5-8рВ-500/О → осушитель
_KM_RCV = re.compile(r"-(\d{3,4})(?:/|\s|$)")             # КМ5.5-8рВ-500 → ресивер
_KM_VSD = re.compile(r"\bЧРП\b", re.I)
_HANS_RCV = re.compile(r"-(\d{3,4})\s*(?:DR|DF|D)?\s*$", re.I)   # RSE 7.5-12-500 DR
_HANS_DRY = re.compile(r"\bDR\b|\bDRY\b", re.I)
_IRON_RCV = re.compile(r"(\d{3,4})\s*L\b", re.I)          # IC 7,5/8 DIGI DF 500L
_IRON_DRY = re.compile(r"\bDF\b|\bDRY\b", re.I)
_EX_TAIL = re.compile(r"\bEX([A-Z]*)", re.I)


def vendor_props(row: dict, specs: dict) -> dict:
    """13 свойств шаблона для карточки сайта производителя.

    Сначала явные характеристики сайта (их публикуют GMP, ET, Exelute, Magnus, Dali,
    CrossAir, ЗИФ, Sollant, Comprag), затем — хвосты кода модели там, где сайт свойства
    не печатает (BERG/ATOM печатают только давление, IP и частотник; Kraftmachine —
    только двигатель, IP и мощность; Hansmann и Ironmac ресивер не печатают вовсе)."""
    brand = row["brand"]
    name = row["name"] or ""
    p = dict.fromkeys(PROPS)

    p["блок"] = block_class(pick(specs, r"винтов\w*\s*блок|марка\s+винтов", exclude=r"гарант"))
    p["ресивер"] = yesno(pick(specs, r"^ресивер$|наличие\s+ресивера|на\s+ресивере"))
    if p["ресивер"] is None:
        v = pick(specs, r"об[ъь]?[её]м\s+ресивера")
        if v is not None:
            p["ресивер"] = "нет" if re.search(r"без|^нет$", v, re.I) else ("да" if num(v) else None)
    p["осушитель"] = yesno_or_model(pick(specs, r"осушител|с\s+осушителем", exclude=r"модель\s+осушител"))
    if p["осушитель"] is None:
        p["осушитель"] = yesno_or_model(pick(specs, r"модель\s+осушител"))
    p["частотник"] = yesno_or_model(pick(specs, r"частотн"))
    if p["частотник"] is None and row.get("vsd"):
        p["частотник"] = yesno(row["vsd"])
    p["фильтры"] = yesno_or_model(pick(specs, r"встроенн\w*\s+фильтр|комплектн\w*\s+фильтр"))
    p["безмасляный"] = yesno(pick(specs, r"^безмасляный$"))
    if p["безмасляный"] is None:
        v = pick(specs, r"тип\s+смазки|масляный/безмасляный")
        if v:
            p["безмасляный"] = "да" if re.search(r"безмасл", v, re.I) else "нет"
    p["охлаждение"] = cooling(pick(specs, r"тип\s+охлаждения|система\s+охлаждения|"
                                          r"тип\s+системы\s+охлаждения"))
    p["ступени"] = stages(pick(specs, r"количество\s+ступеней|число\s+ступеней"))
    if p["ступени"] is None:
        v = yesno(pick(specs, r"^двухступенчатый$"))
        p["ступени"] = {"да": "2", "нет": "1"}.get(v)
    p["двигатель"] = motor_class(pick(specs, r"тип\s+(?:электро)?двигател"))
    if p["двигатель"] is None and yesno(pick(specs, r"двигатель\s+на\s+постоянных\s+магнитах")) == "да":
        p["двигатель"] = "синхронный"
    v = pick(specs, r"^передвижной$|^исполнение$")
    if v:
        p["передвижной"] = "да" if re.search(r"да|шасси|прицеп|гусенич", v, re.I) else \
                           ("нет" if re.search(r"нет|стационар", v, re.I) else None)

    for k, val in from_name(name).items():                # маркеры имени поверх молчания сайта
        if p.get(k) is None:
            p[k] = val

    if brand in ("BERG", "ATOM"):
        code = name.split()[-1]
        if _BERG_RCV.search(code):
            p["ресивер"], p["осушитель"] = "да", ("да" if _BERG_DRY.search(code) else "нет")
        elif p["ресивер"] is None:
            p["ресивер"] = "нет"
    if brand == "KRAFTMACHINE":
        p["осушитель"] = "да" if _KM_DRY.search(name) else (p["осушитель"] or "нет")
        p["ресивер"] = "да" if _KM_RCV.search(name) else (p["ресивер"] or "нет")
        if _KM_VSD.search(name):
            p["частотник"] = "да"
    if brand == "HANSMANN":
        p["ресивер"] = "да" if _HANS_RCV.search(name) else "нет"
        if p["осушитель"] is None:
            p["осушитель"] = "да" if _HANS_DRY.search(name) else None
    if brand == "IRONMAC":
        p["ресивер"] = "да" if _IRON_RCV.search(name) else (p["ресивер"] or "нет")
        if p["осушитель"] is None and _IRON_DRY.search(name):
            p["осушитель"] = "да"
    if brand == "XELERON":
        # Раздел «vintovye-kompressory-dry-tank» и приставка «Dry T250» в имени — это и есть
        # комплектация «ресивер 250 л + осушитель»: у Xeleron она обозначена только так,
        # характеристик ресивера в таблице нет.
        dry = re.search(r"\bdry\s*t\s*\d+", name, re.I) or "dry-tank" in (row.get("url") or "")
        p["ресивер"] = p["осушитель"] = "да" if dry else "нет"
    # Суффикс «на постоянных магнитах» пишут слитно с кодом — Xeleron «Z10PMA»,
    # Kraftmachine «КМ11-8ПМ». \bPM\b на них не срабатывает, поэтому отдельным правилом.
    if p["двигатель"] is None and re.search(r"\d\s*(?:PMA?|ПМ)\b", name):
        p["двигатель"] = "синхронный"

    ip = ip_classes(row.get("ip") or "") or ip_classes(
        pick(specs, r"степень\s+защиты|класс\s+защиты|уровень\s+защиты") or "")
    return {"kw": kw_value(row.get("kw")), "bar": bar_bucket(row.get("bar")), "ip": ip, **p}


_SCREW = re.compile(r"винтов|vintov|screw", re.I)
# Не компрессор вовсе. Слово должно стоять В НАЧАЛЕ названия — это заголовок товара.
# Без привязки к началу regex ловил «маслоЗАПОЛНЕННЫЙ» и выбрасывал 1 262 карточки
# Kraftmachine («Винтовой электрический маслозаполненный компрессор BS15-16ПМ»).
_NOT_COMPR = re.compile(r"^\s*(?:комплектн\w+\s+|рефрижераторн\w+\s+|адсорбционн\w+\s+)?"
                        r"(?:осушител|ресивер|воздухосборник|фильтр|масло|сепаратор|"
                        r"генератор\s+азота|конденсатоотводчик|запчаст|воздуходувк|"
                        r"труб|бустер|дожимн|винтов\w*\s+блок)", re.I)
# Явно другая форма машины. Отсеиваем ТОЛЬКО по явному признаку: у Kraftmachine и Sollant
# винтовые двухступенчатые называются «Двухступенчатый компрессор SLTT-22V» — слова
# «винтовой» нет ни в имени, ни в адресе, и требование этого слова теряло 591 карточку.
_OTHER_FORM = re.compile(r"спиральн|scroll|поршнев|piston|центробежн|мембранн|"
                         r"порш\w*\s+компрессор", re.I)
_NAME_BAR = re.compile(r"\d+(?:[.,]\d+)?\s*/\s*(\d{1,2})(?:\D|$)")


def load_vendors() -> list:
    """Винтовые компрессоры со всех собранных сайтов производителей."""
    out = []
    for path in sorted(VENDORS.glob("*.csv")):
        if path.name.endswith("_sverka.csv"):
            continue
        for row in csv.DictReader(open(path, encoding="utf-8-sig"), delimiter=";"):
            name = row.get("name") or ""
            place = name + " " + (row.get("url") or "")
            # «Винтовой блок Hanbell AC 130» у Sollant и «Осушитель RDX150» у Comprag лежат
            # в тех же разделах, что и компрессоры: режем по имени, а не по разделу.
            if _NOT_COMPR.search(name) or _OTHER_FORM.search(name):
                continue
            if not _SCREW.search(place) and not re.search(r"kompressor|компрессор", place, re.I):
                continue
            if not (row.get("bar") or "").strip():
                # Exelute на 136 карточках не выводит давление в таблицу, но пишет его в
                # названии вторым числом («EXT PM 22/12 IP54» = 22 кВт / 12 бар).
                m = _NAME_BAR.search(name)
                if m:
                    row = dict(row, bar=m.group(1))
            try:
                specs = json.loads(row.get("specs") or "{}")
            except Exception:
                specs = {}
            if not isinstance(specs, dict):
                specs = {}
            p = vendor_props(row, specs)
            if p["kw"] is None or p["bar"] is None:
                continue
            try:
                price = float(row["price"]) if row.get("price") else None
            except ValueError:
                price = None
            out.append(dict(brand=row["brand"], name=name, url=row.get("url"), price=price,
                            file=path.stem, props=p))
    # Дедуп. Одна карточка часто попадает в выгрузку несколько раз: у Kraftmachine
    # «КМ110-10 ВБМ» лежит и в разделе по мощности, и в разделе по давлению, и в одном из
    # них цены нет. Ключ включает весь вектор свойств, поэтому исполнения по IP (у KM это
    # реально разные SKU: IP54 и IP65) остаются раздельными, а полные дубли схлопываются
    # в карточку с ценой — иначе «аналогов 7» на деле означало бы три разных машины.
    best = {}
    for c in out:
        k = (c["brand"], c["name"], tuple(sorted((str(a), str(b)) for a, b in c["props"].items())))
        cur = best.get(k)
        if cur is None or (c["price"] and (not cur["price"] or c["price"] < cur["price"])):
            best[k] = c
    return list(best.values())


# --- сцепка -----------------------------------------------------------------------------
def value(p: dict, key: str, strict: bool):
    """Значение свойства с учётом умолчания.

    strict=True — буква правила заказчика: молчание сайта = первое значение столбца.
    strict=False — умолчание подставляется только когда молчат ОБЕ стороны, то есть
    молчание не режет кандидата (правило 4 «молчание = совместимо»)."""
    v = p.get(key)
    if v is None and strict:
        return DEFAULTS[key]
    return v


def match(o: dict, c: dict, strict: bool) -> tuple[bool, list]:
    """Совпал ли кандидат. Возвращает (да/нет, список свойств, которые держатся на умолчании)."""
    if o["kw"] != c["kw"] or o["bar"] != c["bar"]:
        return False, []
    if o["ip"] and c["ip"] and not (o["ip"] & c["ip"]):     # IP — тристейт, см. enger_specs
        return False, []
    weak = []
    for k in PROPS:
        a, b = value(o, k, strict), value(c, k, strict)
        if a is None or b is None:                          # мягкий режим: молчит хоть одна
            if o.get(k) is None and c.get(k) is None:
                weak.append(k)
            continue
        if a != b:
            return False, []
        if o.get(k) is None or c.get(k) is None:
            weak.append(k)
    return True, weak


def model_label(name: str) -> str:
    """Код модели для ячейки: без слов-категорий, но С сохранением IP — заказчик в образце
    пишет «CA5.5-8RA-300DRY (IP54)», для него исполнение по защите часть модели."""
    s = re.sub(r"(?i)^(винтовой|двухступенчатый|безмасляный|дизельный|электрический|"
               r"маслозаполненный|мобильный|передвижной|компрессор|воздушный|на\s+постоянных\s+"
               r"магнитах|серии|с\s+прямым\s+приводом|с\s+ременным\s+приводом|"
               r"с\s+частотным\s+преобразователем|на\s+ресивере|со\s+встроенным\s+осушителем|"
               r"[,\s])+", " ", name)
    s = re.sub(r"\s+\d{8,}\s*$", "", s)                     # артикул Comprag в хвосте имени
    s = re.sub(r"\s+", " ", s).strip(" ,-·")
    # Comprag пишет комплектацию прозой: «серии F с ременным приводом на ресивере FR-0510-270».
    # Прозу в ячейку тащить незачем — код модели стоит последним и однозначно узнаётся.
    if len(s) > 34:
        m = re.search(r"([A-ZА-Я]{1,5}-?\d[\w./,-]*)\s*$", s)
        if m:
            s = m.group(1)
    return s or name


# Свойства комплектации: ими отличается «машина на ресивере с осушителем» от голой машины.
# Четыре завода (ATOM, Kraftmachine кроме КМ11, Ironmac, Magnus частично) такие исполнения
# на сайте не публикуют вовсе — по строгому ключу у них аналога нет и быть не может.
# Чтобы это не выглядело как «аналога нет в природе», собираем их отдельным листом.
KOMPLEKT = {"ресивер", "осушитель", "фильтры"}


def build(proba=False, strict=True):
    ip_map = our_ip_by_name()
    ours = [o for o in load_ours_all().get("enger", [])
            if _SCREW.search(o.get("name") or "") and o.get("kw") and o.get("bar")]
    if proba:
        ours = [o for o in ours if any(m in o["name"] for m in PROBA)]
    cands = load_vendors()

    by_kw = defaultdict(list)
    for c in cands:
        by_kw[c["props"]["kw"]].append(c)

    rows = []
    for o in ours:
        p = enger_props(o, ip_map)
        hits, soft, base, ladder = (defaultdict(list), defaultdict(list),
                                    defaultdict(list), defaultdict(list))
        ocls = bar_classes(p["bar"])
        for c in by_kw.get(p["kw"], []):
            ok, weak = match(p, c["props"], strict)
            if ok:
                hits[c["brand"]].append((c, weak))
                continue
            # мягкий разбор нужен только там, где строгий ничего не дал — иначе шум
            ok2, weak2 = match(p, c["props"], False)
            if ok2:
                soft[c["brand"]].append((c, weak2))
                continue
            ok3, diff = match_base(p, c["props"])
            if ok3:
                base[c["brand"]].append((c, diff))
                continue
            # последний заход: давление приведено к ряду шаблона с обеих сторон
            if ocls & bar_classes(c["props"]["bar"]) and p["bar"] != c["props"]["bar"]:
                ok4, _ = match(dict(p, bar="~"), dict(c["props"], bar="~"), strict)
                if ok4:
                    ladder[c["brand"]].append((c, sorted(ocls & bar_classes(c["props"]["bar"]))))
        rows.append(dict(o=o, props=p, family=enger_family(o["name"], o.get("bar")),
                         hits=hits, soft=soft, base=base, ladder=ladder))
    return rows, cands


def match_base(o: dict, c: dict) -> tuple[bool, list]:
    """«Та же машина, другая комплектация»: сходится всё, кроме ресивера/осушителя/фильтров."""
    if o["kw"] != c["kw"] or o["bar"] != c["bar"]:
        return False, []
    if o["ip"] and c["ip"] and not (o["ip"] & c["ip"]):
        return False, []
    diff = []
    for k in PROPS:
        a, b = value(o, k, True), value(c, k, True)
        if a == b:
            continue
        if k not in KOMPLEKT:
            return False, []
        diff.append(f"{k}: у нас {a}, у них {b}")
    return bool(diff), diff


_REAL_KEY = re.compile(r"^([A-ZА-Я]{2}-\d+(?:[.,]\d+)?)([BD])([TFS])R?[СC]?(E?)")


def real_flow_table() -> dict:
    """Лист «реальная пр-ть» из образца заказчика: модель → (8 бар, 10 бар), м³/мин.

    Таблица ручная и живёт у заказчика: это его замеры против паспортных цифр. Держим её
    в data/ рядом с проектом, а не читаем из папки загрузок — та живёт одну сессию."""
    out = {}
    path = _HERE / "data" / "enger_realnaya_prt.csv"
    if not path.exists():
        return out
    for r in csv.DictReader(open(path, encoding="utf-8-sig"), delimiter=";"):
        out[(r["модель"] or "").strip()] = (num(r.get("8 бар")), num(r.get("10 бар")))
    return out


def real_key(family: str) -> str:
    """«HC-7,5DFRE-250» → «HC-7,5DFE»: в таблицах заказчика модель записана без ресивера
    («R») и без объёма, но с буквой частотника («E»)."""
    m = _REAL_KEY.match(family)
    return "".join(m.groups()) if m else ""


def report(rows, tag):
    """Числа до/после — без них изменение не считается понятым (правило 2)."""
    n = len(rows)
    with_any = sum(1 for r in rows if r["hits"])
    pairs = sum(len(v) for r in rows for v in r["hits"].values())
    brands = sum(len(r["hits"]) for r in rows)
    print(f"  {tag:<22} строк {n:>5} | с аналогами {with_any:>5} ({with_any/max(n,1):.0%}) | "
          f"пар {pairs:>6} | бренд-совпадений {brands:>5}")
    return with_any, pairs


# --- запись отчёта ----------------------------------------------------------------------
HFILL = PatternFill("solid", fgColor="305496")
HFONT = Font(bold=True, color="FFFFFF")
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
LINK = Font(color="0563C1", underline="single")
RED = PatternFill("solid", fgColor="FFFFC7CE")       # конкурент дешевле нас
GREY = Font(color="808080", italic=True)           # значение держится на умолчании
FIX = ["Модель Enger", "Бар", "Пр-ть, м³/мин", "Реальная\nпр-ть", "Наша цена, ₽",
       "Цена мин", "Цена сред", "Цена new", "% сниж", "Мы дороже на", "Аналогов"]


def head(ws, titles, row=1):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row, i, t)
        c.font, c.fill, c.alignment = HFONT, HFILL, CTR


def cheapest(lst):
    """Самый дешёвый аналог бренда; карточки без цены — в конец (цена нужна для сравнения)."""
    return min(lst, key=lambda t: (t[0]["price"] is None, t[0]["price"] or 0))


# Заливки эталона. Держим их темой книги, а не готовым RGB: у заказчика в файле стоят
# именно theme+tint, и при смене темы Office отчёт должен перекраситься так же, как его
# собственный. theme 7 = accent4 (охра), 8 = accent5 (голубой), 0 = фон.
F_MODEL = PatternFill("solid", start_color=Color(theme=7, tint=0.8))    # колонка модели
F_PRICE = PatternFill("solid", start_color=Color(theme=8, tint=0.6))    # наша цена + шапка
F_CALC = PatternFill("solid", start_color=Color(theme=7, tint=0.8))     # блок расчёта I..N
F_NEW = PatternFill("solid", start_color=Color(theme=9, tint=0.6))      # колонки «new»
F_EMPTY = PatternFill("solid", start_color=Color(theme=0, tint=-0.35))  # аналога нет
F_WHITE = PatternFill("solid", start_color=Color(theme=0, tint=0.0))


def dxf_fill(rgb: str) -> PatternFill:
    """Заливка для правила условного форматирования.

    В обычной ячейке solid-заливка берётся из fgColor, а в dxf (differential formatting,
    то есть в правилах) Excel красит по bgColor. openpyxl из PatternFill("solid",
    start_color=X) пишет только fgColor — правило срабатывало, но фон оставался белым, и
    в отчёте красился один шрифт. У заказчика в образце заданы оба цвета, делаем так же."""
    return PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)

RUB = r'_-* #,##0\ _₽_-;\-* #,##0\ _₽_-;_-* "-"\ _₽_-;_-@_-'
HDR_FONT = Font(bold=True, size=12)
HDR_AL = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Порядок предпочтения давления при выборе строки-представителя модели. Эталон сцепляет
# по 8 бар; если у серии его нет (LUF идёт на 7,5 / 8,6 / 10,4), спускаемся по ряду.
BAR_PREF = ["8", "10", "7", "12", "13", "15/16", "20", "25", "30", "40"]


def svodnaya_etalon(wb, rows, real):
    """Лист «Сводная» в раскладке и цветах присланного образца.

    Строка = МОДЕЛЬ Enger (как у заказчика), на модель отводится две строки — во второй
    идёт следующий по цене аналог бренда; ровно так сделано в образце, где у Cross-air
    в первой строке стоит исполнение IP54, во второй IP23. Колонки A/B/C/H заполняются
    только в первой строке группы, и формулы «разницы» второй строки ссылаются на её H."""
    ws = wb.create_sheet("Сводная", 0)
    fix = ["Модель Enger", "8 бар", "10 бар", "Макс\n пр-ть", "Реальная\nпр-ть", "пр-ть new",
           "% к 10 реальной", "Наша Цена", "цена мин", "цена сред", " цена new", "% сниж",
           "мы дороже на", "накрутка после изменения цены"]
    n_fix = len(fix)
    b0 = n_fix + 1                       # короткий блок «цена по брендам» (O–Z образца)
    t0 = b0 + len(BRANDS)                # тройки «Бренд | Цена | разница» (с AA образца)
    for i, t in enumerate(fix + [TITLE[b] for b in BRANDS] +
                          [x for b in BRANDS for x in (TITLE[b], "Цена", "разница")], 1):
        c = ws.cell(1, i, t)
        c.font, c.alignment = HDR_FONT, HDR_AL
        # В образце голубым залито только имя бренда и «Наша Цена»; «Цена» в тройке белая,
        # «разница» — вообще без заливки. Повторяем ровно так.
        if i in (1, 8) or b0 <= i < t0 or (i >= t0 and (i - t0) % 3 == 0):
            c.fill = F_PRICE
        elif i in (6, 11, 13, 14):       # «пр-ть new», «цена new», «мы дороже», «накрутка»
            c.fill = F_NEW
        elif i >= t0 and (i - t0) % 3 == 2:
            pass
        else:
            c.fill = F_WHITE
    ws.row_dimensions[1].height = 128.25

    groups = defaultdict(dict)
    for r in rows:
        groups[r["family"]].setdefault(r["props"]["bar"], r)
    pcols = [get_column_letter(t0 + 3 * j + 1) for j in range(len(BRANDS))]

    for fam in sorted(groups):
        bars = groups[fam]
        pick_bar = next((b for b in BAR_PREF if b in bars), next(iter(bars)))
        r0 = bars[pick_bar]
        top = ws.max_row + 1
        # сколько строк на модель: как в образце — две, но если у бренда всего один аналог,
        # вторая строка просто останется серой
        depth = min(2, max([len(v) for v in r0["hits"].values()] or [1]))
        for k in range(depth):
            i = top + k
            if k == 0:
                ws.cell(i, 1, fam)
                if r0["o"].get("url"):
                    ws.cell(i, 1).hyperlink = r0["o"]["url"]
                    ws.cell(i, 1).font = Font(color="0563C1", underline="single")
                b8, b10 = bars.get("8"), bars.get("10")
                ws.cell(i, 2, round(b8["o"]["fl"] / 1000, 2) if b8 and b8["o"].get("fl") else None)
                ws.cell(i, 3, round(b10["o"]["fl"] / 1000, 2) if b10 and b10["o"].get("fl") else None)
                rk = real.get(real_key(fam), (None, None))
                ws.cell(i, 5, rk[1] if pick_bar == "10" else rk[0])
                ws.cell(i, 8, r0["o"].get("price"))
                rng = ",".join(f"{c}{i}" for c in pcols)
                # MIN по пустому диапазону даёт 0, а не ошибку — поэтому COUNT, а не IFERROR:
                # без этого «цена мин» показывала 0 ₽ у всех моделей без аналогов
                ws.cell(i, 9, f'=IF(COUNT({rng})=0,"",MIN({rng}))')
                ws.cell(i, 10, f'=IF(COUNT({rng})=0,"",AVERAGE({rng}))')
                ws.cell(i, 12, f'=IF(N(K{i})=0,"",IFERROR((K{i}-H{i})/H{i},""))')
                ws.cell(i, 13, f'=IFERROR((H{i}-I{i})/I{i},"")')
                # Накрутка считается как в образце: цена / закупку. Две поправки к его
                # формуле, обе вынужденные. Первая: ключом идёт не A2, а код без ресивера
                # («BS-5,5BFR-250» → «BS-5,5BF») — в таблице закупок модели записаны так,
                # и VLOOKUP по полному имени возвращал «*» на каждой строке. Вторая: закуп
                # в долларах, а наша цена в рублях, поэтому делим ещё на курс — он лежит
                # одной ячейкой накрутки!$F$1, чтобы его можно было поправить руками.
                rk = real_key(fam) or fam
                buy = f'VLOOKUP("{rk}",накрутки!A:C,2,FALSE)*накрутки!$F$1'
                ws.cell(i, 14, f'=IFERROR(IF(N(K{i})=0,H{i}/({buy}),K{i}/({buy})),"*")')
            for col in (1,):
                ws.cell(i, col).fill = F_MODEL
            for col in range(2, 8):
                ws.cell(i, col).fill = F_WHITE
            ws.cell(i, 8).fill = F_PRICE
            for col in range(9, 15):
                ws.cell(i, col).fill = F_CALC
                ws.cell(i, col).number_format = RUB if col in (9, 10, 11) else (
                    "0%" if col in (12, 13) else "#,##0.00")
            for j, b in enumerate(BRANDS):
                mc, pc, dc = t0 + 3 * j, t0 + 3 * j + 1, t0 + 3 * j + 2
                lst = sorted(r0["hits"].get(b, []),
                             key=lambda t: (t[0]["price"] is None, t[0]["price"] or 0))
                if k == 0:
                    ws.cell(i, b0 + j, f'=IF({get_column_letter(pc)}{i}="","",'
                                       f'{get_column_letter(pc)}{i})')
                    ws.cell(i, b0 + j).number_format = RUB
                if k >= len(lst):
                    for col in (mc, pc, dc):
                        ws.cell(i, col).fill = F_EMPTY      # серым — аналога нет, как в образце
                    continue
                c, weak = lst[k]
                cell = ws.cell(i, mc, model_label(c["name"]))
                if c.get("url"):
                    cell.hyperlink = c["url"]
                    cell.font = Font(color="0563C1", underline="single",
                                     italic=bool(weak))
                elif weak:
                    cell.font = Font(italic=True, color="808080")
                ws.cell(i, pc, c["price"]).number_format = RUB
                ws.cell(i, dc, f'=IFERROR((H{top}-{get_column_letter(pc)}{i})/'
                               f'{get_column_letter(pc)}{i},"")').number_format = "0%"
    # Короткий блок брендов красим ОТНОСИТЕЛЬНО НАШЕЙ ЦЕНЫ, а не сравнением конкурентов
    # между собой. В образце тут стояла трёхцветная шкала по строке (min красный, max
    # зелёный), и на двадцати строках она читалась. На нашем каталоге — нет: из 663 строк
    # с аналогами в 225 аналог всего один (он одновременно min и max, то есть всегда
    # зелёный), ещё в 110 их ровно два (один всегда красный, другой всегда зелёный).
    # В 51% строк цвет не нёс информации. Полярность оставлена та же, что у колонок
    # «разница» в образце: конкурент дешевле нас — красный, дороже — зелёный.
    first = get_column_letter(b0)
    rng = f"{first}2:{get_column_letter(b0 + len(BRANDS) - 1)}{ws.max_row}"
    for op, fill, color in (("<", "FFFFC7CE", "FF9C0006"), (">", "FFC6EFCE", "FF006100")):
        ws.conditional_formatting.add(rng, FormulaRule(
            # обе проверки на пустоту обязательны: у второй строки группы «наша цена»
            # пустая, и без них пустое сравнивалось бы с нулём и красило всю строку
            formula=[f'AND({first}2<>"",$H2<>"",{first}2{op}$H2)'],
            fill=dxf_fill(fill), font=Font(color=color)))
    # Колонки «разница» — то самое «относительно нашей цены». В образце на них висит пара
    # правил Excel «больше 0 → зелёный / меньше 0 → красный»: разница считается как
    # (наша цена − цена конкурента) / цена конкурента, поэтому зелёный = мы дороже,
    # красный = мы дешевле. У заказчика они разложены поячеечно (714 правил, местами с
    # перевёрнутой полярностью от повторного копирования) — кладём по одному правилу на
    # колонку и в той полярности, которая у него преобладает (682 ячейки против 338).
    last = ws.max_row
    for j in range(len(BRANDS)):
        d = get_column_letter(t0 + 3 * j + 2)
        for rule in (CellIsRule(operator="greaterThan", formula=["0"],
                                fill=dxf_fill("FFC6EFCE"), font=Font(color="FF006100")),
                     CellIsRule(operator="lessThan", formula=["0"],
                                fill=dxf_fill("FFFFC7CE"), font=Font(color="FF9C0006"))):
            ws.conditional_formatting.add(f"{d}2:{d}{last}", rule)
    for rule in (CellIsRule(operator="greaterThan", formula=["0"],
                            fill=dxf_fill("FFC6EFCE"), font=Font(color="FF006100")),
                 CellIsRule(operator="lessThan", formula=["0"],
                            fill=dxf_fill("FFFFC7CE"), font=Font(color="FF9C0006"))):
        ws.conditional_formatting.add(f"G2:G{last}", rule)   # «% к 10 реальной», как в образце

    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 27.9
    for col, w in (("B", 9.9), ("C", 8.3), ("D", 1.0), ("E", 14.4), ("F", 8.4), ("G", 13.3),
                   ("H", 12.0), ("I", 16.4), ("J", 11.4), ("K", 11.4), ("L", 1.0),
                   ("M", 11.9), ("N", 13.0)):
        ws.column_dimensions[col].width = w
    for j in range(len(BRANDS)):
        ws.column_dimensions[get_column_letter(b0 + j)].width = 11.9
        ws.column_dimensions[get_column_letter(t0 + 3 * j)].width = 26
        ws.column_dimensions[get_column_letter(t0 + 3 * j + 1)].width = 12.3
        ws.column_dimensions[get_column_letter(t0 + 3 * j + 2)].width = 9.9
    return ws


def copy_table(wb, title, path, widths=(26, 14, 14)):
    """Таблицы заказчика («накрутки», «реальная пр-ть») кладём в книгу как отдельные листы —
    иначе VLOOKUP в колонке «накрутка» ссылается в пустоту."""
    ws = wb.create_sheet(title)
    if not Path(path).exists():
        return ws
    for r in csv.reader(open(path, encoding="utf-8-sig"), delimiter=";"):
        ws.append([num(v) if i and v not in ("", "*") else v for i, v in enumerate(r)])
    for i in range(1, 4):
        c = ws.cell(1, i)
        c.font, c.fill, c.alignment = HDR_FONT, F_PRICE, HDR_AL
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    if title == "накрутки":
        # Курс отдельной ячейкой, а не константой в формуле: закуп в таблице заказчика
        # в долларах, наши цены в рублях, и курс он правит чаще, чем всё остальное.
        ws["E1"], ws["F1"] = "Курс ₽/$", 78.9
        ws["E1"].font = HDR_FONT
        ws.column_dimensions["E"].width = 12
    ws.freeze_panes = "A2"
    return ws


def save(rows, cands, out=OUT):
    real = real_flow_table()
    wb = openpyxl.Workbook()

    # --- ЛИСТ «Сводная по барам» — строка на каждое давление --------------------------
    ws = wb.active
    ws.title = "Сводная по барам"
    n_fix = len(FIX)
    b0 = n_fix + 1                       # блок «цена по брендам» (как колонки O–Z образца)
    t0 = b0 + len(BRANDS)                # тройки «Модель | Цена | разница»
    head(ws, FIX + [TITLE[b] for b in BRANDS] +
             [x for b in BRANDS for x in (TITLE[b], "Цена", "разница")])
    for r in rows:
        o, p = r["o"], r["props"]
        i = ws.max_row + 1
        rk = real.get(real_key(r["family"]), (None, None))
        rf = rk[0] if p["bar"] == "8" else (rk[1] if p["bar"] == "10" else None)
        ws.cell(i, 1, r["family"])
        if o.get("url"):
            ws.cell(i, 1).hyperlink = o["url"]
            ws.cell(i, 1).font = LINK
        ws.cell(i, 2, num(p["bar"]) if p["bar"] != "15/16" else "15/16")
        ws.cell(i, 3, round(o["fl"] / 1000, 3) if o.get("fl") else None)
        ws.cell(i, 4, rf)
        ws.cell(i, 5, o.get("price"))
        cols = [get_column_letter(t0 + 3 * j + 1) for j in range(len(BRANDS))]
        rng = ",".join(f"{c}{i}" for c in cols)
        ws.cell(i, 6, f"=IFERROR(MIN({rng}),\"\")")
        ws.cell(i, 7, f"=IFERROR(AVERAGE({rng}),\"\")")
        ws.cell(i, 9, f'=IF(N(H{i})=0,"",IFERROR((H{i}-E{i})/E{i},""))')
        ws.cell(i, 10, f'=IFERROR((E{i}-F{i})/F{i},"")')
        ws.cell(i, 11, sum(len(v) for v in r["hits"].values()))
        for j, b in enumerate(BRANDS):
            lst = r["hits"].get(b)
            mc, pc, dc = t0 + 3 * j, t0 + 3 * j + 1, t0 + 3 * j + 2
            ws.cell(i, b0 + j, f"={get_column_letter(pc)}{i}")
            if not lst:
                continue
            c, weak = cheapest(lst)
            label = model_label(c["name"])
            if len(lst) > 1:
                label += f"  (+{len(lst)-1})"          # остальные — на листе «Все совпадения»
            cell = ws.cell(i, mc, label)
            if c.get("url"):
                cell.hyperlink, cell.font = c["url"], LINK
            if weak:
                cell.font = GREY                        # сцепка держится на умолчании шаблона
            ws.cell(i, pc, c["price"])
            ws.cell(i, dc, f'=IFERROR((E{i}-{get_column_letter(pc)}{i})/'
                           f'{get_column_letter(pc)}{i},"")')
            if c["price"] and o.get("price") and c["price"] < o["price"]:
                ws.cell(i, pc).fill = RED
    for i in range(2, ws.max_row + 1):
        for col in [5, 6, 7, 8] + [t0 + 3 * j + 1 for j in range(len(BRANDS))] + \
                   [b0 + j for j in range(len(BRANDS))]:
            ws.cell(i, col).number_format = "# ##0"
        for col in [9, 10] + [t0 + 3 * j + 2 for j in range(len(BRANDS))]:
            ws.cell(i, col).number_format = "0%"
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 26
    for j in range(len(BRANDS)):
        ws.column_dimensions[get_column_letter(t0 + 3 * j)].width = 30
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    # --- ЛИСТ «Свойства Enger» — из чего сложился ключ --------------------------------
    ws2 = wb.create_sheet("Свойства Enger")
    head(ws2, ["Модель Enger", "Бар", "кВт", "IP"] + PROPS + ["взято из выгрузки"])
    for r in rows:
        p = r["props"]
        i = ws2.max_row + 1
        ws2.append([r["family"], p["bar"], p["kw"],
                    "/".join(sorted(p["ip"])) if p["ip"] else ""] +
                   [p.get(k) or DEFAULTS[k] for k in PROPS] +
                   [sum(1 for k in PROPS if p.get(k) is not None)])
        for j, k in enumerate(PROPS, 5):
            if p.get(k) is None:
                ws2.cell(i, j).font = GREY              # серым — то, что подставлено умолчанием
    ws2.freeze_panes = "B2"
    ws2.column_dimensions["A"].width = 26

    # --- ЛИСТ «Все совпадения» --------------------------------------------------------
    ws3 = wb.create_sheet("Все совпадения")
    head(ws3, ["Модель Enger", "Бар", "Наша цена", "Бренд", "Модель конкурента", "Цена",
               "Разница", "На умолчании шаблона", "Ссылка"])
    for r in rows:
        for b in BRANDS:
            for c, weak in sorted(r["hits"].get(b, []),
                                  key=lambda t: (t[0]["price"] is None, t[0]["price"] or 0)):
                i = ws3.max_row + 1
                ws3.append([r["family"], r["props"]["bar"], r["o"].get("price"), TITLE[b],
                            model_label(c["name"]), c["price"],
                            (r["o"]["price"] - c["price"]) / c["price"]
                            if c["price"] and r["o"].get("price") else None,
                            ", ".join(weak), c.get("url")])
                ws3.cell(i, 9).font = LINK
                if c.get("url"):
                    ws3.cell(i, 9).hyperlink = c["url"]
                ws3.cell(i, 7).number_format = "0%"
    ws3.freeze_panes = "A2"
    for col, w in (("A", 26), ("E", 34), ("H", 30), ("I", 40)):
        ws3.column_dimensions[col].width = w

    # --- ЛИСТ «Другая комплектация» ---------------------------------------------------
    ws4 = wb.create_sheet("Другая комплектация")
    head(ws4, ["Модель Enger", "Бар", "Наша цена", "Бренд", "Модель конкурента", "Цена",
               "Чем отличается", "Ссылка"])
    for r in rows:
        for b in BRANDS:
            if r["hits"].get(b):        # строгий аналог у бренда есть — базовая версия не нужна
                continue
            for c, diff in sorted(r["base"].get(b, []),
                                  key=lambda t: (t[0]["price"] is None, t[0]["price"] or 0))[:2]:
                i = ws4.max_row + 1
                ws4.append([r["family"], r["props"]["bar"], r["o"].get("price"), TITLE[b],
                            model_label(c["name"]), c["price"], "; ".join(diff), c.get("url")])
                if c.get("url"):
                    ws4.cell(i, 8).hyperlink, ws4.cell(i, 8).font = c["url"], LINK
    ws4.freeze_panes = "A2"
    for col, w in (("A", 26), ("E", 34), ("G", 46), ("H", 40)):
        ws4.column_dimensions[col].width = w

    # --- ЛИСТ «Аналог по классу давления» ----------------------------------------------
    # Отдельно, а не в «Сводной»: тут давление приведено к ряду шаблона (наши 8,6 бар против
    # их 8), то есть допуск, которого заказчик просил не делать. Без этого листа 323 строки
    # безмасляных серий не сцеплялись вовсе — паспортных 7,5/8,5/8,6/12,5 бар в ряду нет.
    ws45 = wb.create_sheet("Аналог по классу давления")
    head(ws45, ["Модель Enger", "Наш бар", "Класс ряда", "Наша цена", "Бренд",
                "Модель конкурента", "Их бар", "Цена", "Разница", "Ссылка"])
    for r in rows:
        for b in BRANDS:
            if r["hits"].get(b) or r["soft"].get(b) or r["base"].get(b):
                continue
            for c, cls in sorted(r["ladder"].get(b, []),
                                 key=lambda t: (t[0]["price"] is None, t[0]["price"] or 0))[:2]:
                i = ws45.max_row + 1
                ws45.append([r["family"], r["props"]["bar"], "/".join(cls), r["o"].get("price"),
                             TITLE[b], model_label(c["name"]), c["props"]["bar"], c["price"],
                             (r["o"]["price"] - c["price"]) / c["price"]
                             if c["price"] and r["o"].get("price") else None, c.get("url")])
                ws45.cell(i, 9).number_format = "0%"
                if c.get("url"):
                    ws45.cell(i, 10).hyperlink, ws45.cell(i, 10).font = c["url"], LINK
    ws45.freeze_panes = "A2"
    for col, w in (("A", 26), ("F", 34), ("J", 40)):
        ws45.column_dimensions[col].width = w

    # --- ЛИСТ «Кандидаты (сайт молчит)» ------------------------------------------------
    ws5 = wb.create_sheet("Кандидаты (сайт молчит)")
    head(ws5, ["Модель Enger", "Бар", "Наша цена", "Бренд", "Модель конкурента", "Цена",
               "Молчат свойства", "Ссылка"])
    for r in rows:
        for b in BRANDS:
            if r["hits"].get(b):
                continue
            for c, weak in sorted(r["soft"].get(b, []),
                                  key=lambda t: (t[0]["price"] is None, t[0]["price"] or 0))[:2]:
                i = ws5.max_row + 1
                ws5.append([r["family"], r["props"]["bar"], r["o"].get("price"), TITLE[b],
                            model_label(c["name"]), c["price"], ", ".join(weak), c.get("url")])
                if c.get("url"):
                    ws5.cell(i, 8).hyperlink, ws5.cell(i, 8).font = c["url"], LINK
    ws5.freeze_panes = "A2"
    for col, w in (("A", 26), ("E", 34), ("G", 46), ("H", 40)):
        ws5.column_dimensions[col].width = w

    # --- ЛИСТ «Покрытие и правила» -----------------------------------------------------
    ws6 = wb.create_sheet("Покрытие и правила")
    ws6.column_dimensions["A"].width = 22
    ws6.column_dimensions["B"].width = 104
    by_brand = Counter()
    priced = Counter()
    for c in cands:
        by_brand[c["brand"]] += 1
        priced[c["brand"]] += bool(c["price"])
    hit_brand = Counter()
    for r in rows:
        for b in r["hits"]:
            hit_brand[b] += 1
    n = len(rows)
    cov = Counter()
    for r in rows:
        if r["hits"]:
            cov["строгий"] += 1
        elif r["base"]:
            cov["комплектация"] += 1
        elif r["soft"]:
            cov["молчит"] += 1
        elif r["ladder"]:
            cov["класс бар"] += 1
        else:
            cov["ничего"] += 1
    lines = [
        ("ПОКРЫТИЕ", f"строк (карточек Enger) всего {n}"),
        ("  строгий аналог", f'{cov["строгий"]} ({cov["строгий"]/n:.0%}) — лист «Сводная»'),
        ("  др. комплектация", f'{cov["комплектация"]} — только на листе «Другая комплектация»'),
        ("  сайт молчит", f'{cov["молчит"]} — только на листе «Кандидаты (сайт молчит)»'),
        ("  класс давления", f'{cov["класс бар"]} — только на листе «Аналог по классу давления»'),
        ("  кандидата нет", f'{cov["ничего"]} ({cov["ничего"]/n:.0%}) — у собранных заводов такой '
                            f"машины нет вовсе. Это почти целиком безмасляные серии сухого "
                            f"сжатия (OF, OFS, OFSA, OFSZ) на 3–5 и 20–40 бар."),
        ("", ""),
        ("Правило сцепки", "13 свойств шаблона, точное совпадение. кВт и бар обязательны "
                           "с обеих сторон. Десять свойств, выделенных в шаблоне зелёным, при "
                           "молчании сайта заменяются первым значением столбца."),
        ("Степень защиты", "в шаблоне зелёным НЕ выделена, поэтому умолчание IP23 не "
                           "подставляем: у Hansmann и Xeleron IP не указан ни на одной карточке, "
                           "и подстановка отсекла бы бренд целиком. IP режет пару только при "
                           "явном конфликте обеих сторон."),
        ("Строка отчёта", "карточка Enger = модель + давление. Схлопывать бары нельзя: "
                          "BS-5,5BFR-250 на 8 бар стоит 312 125 ₽, на остальных 327 844 ₽."),
        ("Цены", "рубли, с сайтов производителей. В образце заказчика цены в долларах по "
                 "курсу 78,9 (лист «Лист1» образца: =A1/78.9)."),
        ("ЦВЕТА «Сводной»", "три независимых слоя, их легко перепутать:"),
        ("  заливки блоков", "охра — колонка модели и блок расчёта, голубой — наша цена, "
                             "серый в тройке — аналога у бренда нет. От чисел не зависят."),
        ("  цены брендов", "короткий блок колонок: цена конкурента НИЖЕ нашей — красный, "
                           "ВЫШЕ — зелёный. В образце тут стояла шкала «дешевле всех из "
                           "конкурентов — красный», но на полном каталоге она не читается: "
                           "в 225 строках аналог всего один и всегда красился зелёным, ещё в "
                           "110 их ровно два — один всегда красный, другой зелёный."),
        ("  колонки «разница»", "как в образце: больше нуля — зелёный (мы дороже), меньше — "
                                "красный (мы дешевле)."),
        ("Курсив в тройке", "модель сцеплена с опорой на умолчание шаблона, а не на данные сайта."),
        ("", ""),
        ("Серым в «Свойствах»", "значение подставлено умолчанием шаблона, а не прочитано с сайта."),
        ("Класс давления", "паспортных 7,5 / 8,5 / 8,6 / 12,5 бар (наши безмасляные LUF, OFSA, "
                           "OFS) в ряду шаблона нет, а у заводов те же машины стоят на 7/8/10/12. "
                           "Такие пары вынесены на отдельный лист: в «Сводной» допусков нет."),
        ("", ""),
        ("Remeza", "remeza.com — сайт завода, цен нет ни в разделах, ни в карточках "
                   "(проверено 24.08). Колонка оставлена пустой: брать цену с площадки нельзя."),
        ("Xeleron", "xelerone.com публикует характеристики 76 моделей, но карточки с ценой "
                    "только три (Z7.5A, Z10A, Z175A) — остальное «цена по запросу»."),
        ("Комплектации", "ATOM, Ironmac и Kraftmachine (кроме КМ11) исполнений «на ресивере "
                         "с осушителем» на сайте не публикуют вовсе — по строгому ключу "
                         "аналога у них нет. Их базовые машины — лист «Другая комплектация»."),
        ("", ""),
        ("Бренд", "карточек с сайта / из них с ценой / наших строк, где нашёлся аналог"),
    ]
    for k, v in lines:
        ws6.append([k, v])
    for b in BRANDS:
        ws6.append([TITLE[b], f"{by_brand.get(b,0)} / {priced.get(b,0)} / {hit_brand.get(b,0)}"])
    for i in range(1, ws6.max_row + 1):
        ws6.cell(i, 1).font = Font(bold=True)
        ws6.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")

    # Таблицы заказчика — чтобы формула «накрутка» и колонка «Реальная пр-ть» имели источник
    copy_table(wb, "накрутки", _HERE / "data" / "enger_nakrutki.csv")
    copy_table(wb, "реальная пр-ть", _HERE / "data" / "enger_realnaya_prt.csv")
    svodnaya_etalon(wb, rows, real)            # первый лист — раскладка эталона
    wb.save(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--proba", action="store_true", help="только 10 моделей из образца")
    ap.add_argument("--sravnit", action="store_true", help="сравнить два режима умолчаний")
    a = ap.parse_args()
    if a.sravnit:
        for strict in (True, False):
            report(build(proba=a.proba, strict=strict)[0],
                   "строгое умолчание" if strict else "молчание=совместимо")
        raise SystemExit(0)
    rows, cands = build(proba=a.proba, strict=True)
    report(rows, "строгое умолчание")
    path = save(rows, cands)
    print(f"-> {path}")
